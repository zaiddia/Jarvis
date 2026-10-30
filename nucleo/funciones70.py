"""Lo de la 7.0: tarjetas, diario, proactividad, vigilante, habilidades, MCP, celular y tablero.

Parte de Jarvis (antes dentro de jarvis.py). Los nombres globales de Jarvis se usan como J.<nombre>.
"""
from collections import deque
from flask import Flask
from flask import abort
from flask import render_template
from flask import request
import aula
import datetime
import diario
import habilidades
import hmac
import infodia
import json
import mcp_cliente
import os
import proactivo
import rapidez
import re
import rutinas
import skills
import socket
import sys
import threading
import time
import uuid
import vigilante
import webbrowser

from nucleo import J


# =============================================================================
# 7.0 · TARJETAS DEL HUD: lo que Jarvis cuenta, también a la vista (clima, agenda, rutinas, avisos…)
# =============================================================================
_CARD_ORIGINS = {
    "CLIMA": ("clima", "☁", "Clima"), "AGENDA": ("agenda", "▦", "Agenda"), "NOTAS": ("notas", "✎", "Notas"),
    "CALENDARIO": ("agenda", "▦", "Próximo"), "TIMER": ("timer", "⏱", "Temporizador"), "ALARMA": ("timer", "⏰", "Recordatorio"),
    "RUTINA": ("rutina", "▶", "Rutina"), "PRUEBA": ("prueba", "◎", "Prueba completa"), "MEMORIA": ("memoria", "◈", "Memoria"),
    "PROACTIVO": ("proactivo", "✦", "Jarvis"), "VIGILANTE": ("vigilante", "◉", "Aviso"), "HABILIDAD": ("habilidad", "✧", "Habilidad"),
    "DIARIO": ("diario", "☰", "Diario"), "MCP": ("mcp", "⧉", "Conector"), "TABLERO": ("tablero", "▣", "Tablero"),
    "NOTICIAS": ("noticias", "≡", "Noticias"), "ACTUALIZACIÓN": ("update", "⟳", "Actualización"), "MATE": ("mate", "∑", "Cálculo"),
}


def card(kind, title, body="", icon="✦", secs=9, steps=None, step=None):
    """Tarjeta holográfica en el HUD (y en el panel del celular)."""
    c = {"kind": kind, "title": str(title)[:60], "body": str(body)[:400], "icon": icon, "secs": secs,
         "time": datetime.datetime.now().strftime("%H:%M")}
    if steps is not None:
        c.update(steps=[str(s)[:60] for s in steps][:12], step=step)
    J.recent_cards.append(c)
    try:
        J.socketio.emit("card", c)
    except Exception:
        pass


recent_cards = deque(maxlen=12)


def _auto_card(text, origin):
    k = J._CARD_ORIGINS.get(origin)
    if k and J.settings.get("hud_cards", True) and len(str(text)) > 3:
        J.card(k[0], k[2], text, k[1], secs=min(16, 6 + len(str(text)) // 40))


# =============================================================================
# 7.0 · DIARIO (memoria del día a día): «¿qué hice ayer?», «resume mi semana»
# =============================================================================
_DIARY_RE = re.compile(r"\b(?:que hice|que hicimos|que he hecho|resum\w* (?:de )?(?:mi |la |el )?(?:dia|semana|jornada)|"
                       r"resumen de (?:hoy|ayer|la semana|mi semana|mi dia)|como fue (?:mi|el) dia|mi diario)\b")


_SPEED_RE = re.compile(r"\b(?:que tan rapido|cuan rapido|que tan lento|cuanto tardas|cuanto te demoras|"
                       r"por que tardas|por que te demoras|por que eres tan lento|tiempo de respuesta|"
                       r"velocidad de respuesta|eres lento)\b")


def speed_voice(text, t):
    if not J._SPEED_RE.search(t):
        return False
    J.reply(rapidez.spoken(rapidez.summary()), "TIEMPOS")
    return True


def diary_voice(text, t):
    if re.search(r"\b(?:borra|borrar|elimina|vacia|limpia)\b.*\bdiario\b", t):
        diario.clear()
        J.reply("Listo, borré tu diario. Desde ahora empiezo uno nuevo.", "DIARIO")
        return True
    if not J._DIARY_RE.search(t):
        return False
    if not J.settings["diary_enabled"]:
        J.reply("El diario está apagado. Lo enciendes en Ajustes, sección Memoria.", "DIARIO")
        return True
    today = datetime.date.today()
    if re.search(r"\bsemana\b", t):
        d1, label = today - datetime.timedelta(days=6), "esta semana"
    elif re.search(r"\bantier|anteayer\b", t):
        d1 = today - datetime.timedelta(days=2)
        label = "anteayer"
        today = d1
    elif re.search(r"\bayer\b", t):
        d1 = today = today - datetime.timedelta(days=1)
        label = "ayer"
    else:
        d1, label = today, "hoy"
    entries = diario.days(d1, today)
    if not entries:
        J.reply(f"No tengo nada anotado de {label}.", "DIARIO")
        return True
    try:
        if J.state.get("ai") != "ok":
            raise J.AIError("sin IA", kind="config")
        out = J.ai_chat([{"role": "system", "content": J.settings["system_prompt"] + J.VOICE_RULE + J._user_ctx()},
                       {"role": "user", "content": diario.prompt(entries, label)}], max_tokens=500)
        J.reply(out.strip() or diario.spoken(entries, label), "DIARIO")
    except Exception:
        J.reply(diario.spoken(entries, label), "DIARIO")
    return True


# =============================================================================
# 7.0 · PROACTIVIDAD: Jarvis habla primero (con criterio). Lógica en proactivo.py.
# =============================================================================
pro_state = proactivo.State()
pro_offer = {"routine": None, "expires": 0.0}
_pro_weather = {"day": None, "hourly": []}


def _pro_cfg():
    return {k: J.settings[k] for k in J.settings if k.startswith("proactive_")}


def _pro_say(key, text, optional=True):
    now = datetime.datetime.now()
    if not proactivo.can_speak(J.pro_state, now, J._pro_cfg(), optional):
        return False
    proactivo.said(J.pro_state, now, key, optional)
    J.hud_show()
    J.reply(text, "PROACTIVO")
    diario.log("aviso", text)
    if J.settings["telegram_alerts"] and J.settings["telegram_remote"] and J._idle_seconds() > 600:
        threading.Thread(target=J.telegram_send, args=(f"✦ {text}",), daemon=True).start()
    return True


def offer_routine_for(event_title):
    """Un evento empieza pronto: «¿activo el modo clase?» (y un «sí» la lanza)."""
    if not J.settings["proactive_offer_routines"]:
        return
    r = proactivo.routine_for_event(event_title, rutinas.routines)
    if not r or J.routine_state["running"]:
        return
    if J._pro_say(f"ofrece:{event_title}", f"¿Activo la rutina {r['nombre']}?", optional=False):
        J.pro_offer.update(routine=r, expires=time.time() + 180)


def proactive_answer(t):
    """Respuesta a «¿activo la rutina…?». True si la contestó."""
    r = J.pro_offer["routine"]
    if not r or time.time() > J.pro_offer["expires"]:
        return False
    if re.fullmatch(rf"{J.CONFIRM_RE}(?: .*)?", t) or re.fullmatch(r"(?:activala|lanzala|si activala|hazlo|ponla)", t):
        J.pro_offer.update(routine=None, expires=0)
        J.run_routine(r)
        return True
    if re.fullmatch(rf"{J.DENY_RE}(?: .*)?", t):
        J.pro_offer.update(routine=None, expires=0)
        J.reply("Vale, la dejo para después.", "PROACTIVO")
        return True
    return False


def _pro_hourly():
    """Probabilidad de lluvia por hora de hoy (una consulta al día)."""
    today = datetime.date.today()
    if J._pro_weather["day"] == today:
        return J._pro_weather["hourly"]
    J._pro_weather.update(day=today, hourly=[])
    city = J.settings["weather_city"].strip()
    if not city or J.is_offline():
        return []
    try:
        g = J._geo(city)
        if g:
            url = infodia.forecast_url(g["lat"], g["lon"]).replace(infodia.FORECAST_URL, J.FORECAST_URL)
            h = (json.loads(J._get_text(url)).get("hourly") or {})
            J._pro_weather["hourly"] = list(zip(h.get("time") or [], h.get("precipitation_probability") or []))
    except Exception as e:
        J.log_terminal(f"Avisos (lluvia): {e}", "AVISO")
    return J._pro_weather["hourly"]


def smart_greeting():
    """El saludo del primer uso del día: la hora, el clima en una frase y lo que tienes hoy."""
    now = datetime.datetime.now()
    name = J.settings["address_name"].strip()
    parts = [f"Buenos días{', ' + name if name else ''}. Son las {now:%H:%M}."]
    if J.settings["weather_city"].strip() and not J.is_offline():
        w = J.weather_text()
        if w and not w.startswith(("No ", "Para ", "Ahora")):
            parts.append(re.split(r"(?<=\.)\s", w)[0])
    try:
        items = J.agenda("hoy", with_day=False)
    except Exception:
        items = []
    if items:
        parts.append(f"Hoy tienes {len(items)} cosa{'s' if len(items) != 1 else ''} en la agenda; la primera: {items[0]}.")
    return " ".join(parts)


def _proactive_worker():
    time.sleep(60)
    while True:
        try:
            now = datetime.datetime.now()
            cfg = J._pro_cfg()
            idle = J._idle_seconds() if J.IS_WIN else 0.0
            proactivo.activity(J.pro_state, now, idle)
            if cfg.get("proactive_enabled", True):
                if proactivo.greeting_due(J.pro_state, now, idle, cfg, done_today=J._upd_load().get("saludo") == now.date().isoformat()):
                    if J._pro_say("saludo", J.smart_greeting()):
                        J._upd_save(saludo=now.date().isoformat())
                for due in (proactivo.break_due(J.pro_state, now, cfg),):
                    if due:
                        J._pro_say(*due)
                bat = None
                try:
                    b = J.psutil.sensors_battery() if J.psutil else None
                    bat = (b.percent, b.power_plugged) if b else None
                except Exception:
                    pass
                if bat:
                    due = proactivo.battery_due(J.pro_state, bat[0], bat[1], cfg)
                    if due:
                        J._pro_say(*due)
                if J.settings["leave_time"].strip():
                    due = proactivo.rain_due(J.pro_state, now, J.settings["leave_time"], J._pro_hourly(), cfg)
                    if due:
                        J._pro_say(*due)
        except Exception as e:
            J.log_terminal(f"Avisos proactivos: {e}", "AVISO")
        time.sleep(30)


# =============================================================================
# 7.0 · «AVÍSAME CUANDO TERMINE…» (lógica en vigilante.py)
# =============================================================================
watches = []
psutil = skills.psutil                     # opcional (puede faltar)


def _downloads_dir():
    return os.path.join(os.path.expanduser("~"), "Downloads")


def _partial_downloads():
    try:
        return [f for f in os.listdir(J._downloads_dir()) if f.lower().endswith(vigilante.PARTIAL)]
    except OSError:
        return []


def watch_voice(text, t):
    if re.search(r"\b(?:que (?:estas )?vigilando|que avisos tienes pendientes)\b", t):
        if not J.watches:
            J.reply("No estoy vigilando nada ahora mismo.", "VIGILANTE")
        else:
            J.reply("Estoy pendiente de: " + ", ".join(w.label for w, _ in J.watches) + ".", "VIGILANTE")
        return True
    if re.search(r"\b(?:deja de vigilar|ya no me avises|cancela (?:los |el )?avisos? de)\b", t):
        n = len(J.watches)
        J.watches.clear()
        J.reply("Listo, dejo de vigilar." if n else "No estaba vigilando nada.", "VIGILANTE")
        return True
    req = vigilante.parse(text)
    if not req:
        return False
    if J.psutil is None and req["kind"] != "descarga":
        J.reply("Para vigilar programas me falta un componente (psutil). Vuelve a ejecutar instalar.bat.", "VIGILANTE")
        return True
    if req["kind"] == "descarga" and not os.path.isdir(J._downloads_dir()):
        J.reply("No encuentro tu carpeta de Descargas.", "VIGILANTE")
        return True
    J.watches.append((vigilante.Watch(req), {"remote": bool(getattr(J._tl, "remote", False))}))
    J.reply(f"Vale, te aviso cuando termine {req['label']}." if req["kind"] != "cierre"
          else f"Vale, te aviso cuando se cierre {req['target']}.", "VIGILANTE")
    return True


def _watch_worker():
    cpu_prev = {}
    while True:
        time.sleep(5)
        if not J.watches:
            continue
        try:
            now = datetime.datetime.now()
            partial = J._partial_downloads()
            running, cpu_of = [], {}
            if J.psutil is not None and any(w.kind != "descarga" for w, _ in J.watches):
                for p in J.psutil.process_iter(["name", "pid"]):
                    try:
                        name = (p.info["name"] or "").lower().replace(".exe", "")
                        running.append(name)
                        if any(name and vigilante._match(vigilante.procs_for(w.target), [name]) for w, _ in J.watches if w.kind == "proceso"):
                            proc = cpu_prev.get(p.info["pid"]) or p
                            cpu_prev[p.info["pid"]] = proc
                            cpu_of[name] = cpu_of.get(name, 0.0) + proc.cpu_percent(None) / max(1, J.psutil.cpu_count() or 1)
                    except (J.psutil.NoSuchProcess, J.psutil.AccessDenied):
                        continue
            for item in list(J.watches):
                w, meta = item
                cpu = None
                if w.kind == "proceso":
                    hits = vigilante._match(vigilante.procs_for(w.target), list(cpu_of))
                    cpu = sum(cpu_of[h] for h in hits) if hits else None
                out = w.step(now, partial, running, cpu)
                if out is None:
                    continue
                J.watches.remove(item)
                if out == "timeout":
                    J.reply(f"No vi {w.label} en marcha, así que dejo de vigilarla." if w.kind == "descarga"
                          else f"Dejo de vigilar {w.label}: pasaron demasiadas horas.", "VIGILANTE")
                elif out == "notfound":
                    J.reply(f"No veo {w.target} abierto, así que no hay nada que vigilar.", "VIGILANTE")
                else:
                    J.hud_show()
                    J.play_sound(J.resolve_alarm_sound(), max_seconds=3)
                    J.alert(out, "VIGILANTE")
                    diario.log("aviso", out)
                    if meta.get("remote") and not (J.settings["telegram_alerts"] and J.settings["telegram_remote"]):
                        threading.Thread(target=J.telegram_send, args=(f"🔔 {out}",), daemon=True).start()
        except Exception as e:
            J.log_terminal(f"Vigilante: {e}", "AVISO")


# =============================================================================
# 7.0 · HABILIDADES INSTALABLES (lógica en habilidades.py)
# =============================================================================
habilidades.configure(J.CONFIG_DIR)


def skill_voice(text, t):
    if re.search(r"\b(?:que|cuales) habilidades\b|\bmis habilidades\b|\bhabilidades instaladas\b", t):
        hs = habilidades.skills
        if not hs:
            J.reply("No tienes habilidades instaladas. Se instalan en Ajustes, sección Rutinas.", "HABILIDAD")
        else:
            J.reply(f"Tienes {len(hs)} habilidades: " + "; ".join(f"{h['nombre']} («{h['frases'][0]}»)" for h in hs[:8]) + ".", "HABILIDAD")
        return True
    m = habilidades.match(text)
    if not m:
        return False
    h, vals = m
    diario.log("habilidad", h["nombre"])
    J.log_terminal(f"Habilidad «{h['nombre']}» {vals}", "HABILIDAD")
    if h["accion"] == "abrir":
        url = habilidades.fill(h["url"], vals, url=True)
        webbrowser.open(url)
        J.reply(f"Abriendo {h['nombre']}.", "HABILIDAD")
    elif h["accion"] == "decir":
        J.reply(habilidades.fill(h["texto"], vals), "HABILIDAD")
    elif h["accion"] == "ia":
        prompt = habilidades.fill(h["texto"], vals)
        if getattr(J._tl, "remote", False) or getattr(J._tl, "capture", None) is not None:
            J.reply(J.ask_ai(prompt)[1], "HABILIDAD")
        else:
            J.answer_question_streaming(prompt)
    else:
        J.run_routine({"nombre": h["nombre"], "pasos": [habilidades.fill(p, vals) for p in h["pasos"]]})
    return True


@J.socketio.on("habilidades_get")
def handle_habilidades_get(data=None):
    errors = habilidades.load()
    return {"habilidades": habilidades.snapshot(), "errores": [f"{a}: {b}" for a, b in errors], "carpeta": habilidades._dir}


@J.socketio.on("habilidades_install")
def handle_habilidades_install(data):
    try:
        h = habilidades.install(str((data or {}).get("texto", "")))
    except ValueError as e:
        return {"ok": False, "error": str(e), "habilidades": habilidades.snapshot()}
    J.log_terminal(f"Habilidad instalada: {h['nombre']}", "HABILIDAD")
    return {"ok": True, "nombre": h["nombre"], "habilidades": habilidades.snapshot()}


@J.socketio.on("habilidades_delete")
def handle_habilidades_delete(data):
    habilidades.delete(str((data or {}).get("nombre", "")))
    return {"habilidades": habilidades.snapshot()}


@J.socketio.on("habilidades_export")
def handle_habilidades_export(data):
    return {"texto": habilidades.export(str((data or {}).get("nombre", ""))) or ""}


@J.socketio.on("habilidades_folder")
def handle_habilidades_folder(data=None):
    if J.IS_WIN:
        try:
            os.startfile(habilidades._dir)
        except OSError:
            pass
    return {"ok": True}


# =============================================================================
# 7.0 · CONECTORES MCP (lógica en mcp_cliente.py) y Jarvis como servidor MCP (jarvis_mcp.py)
# =============================================================================
mcp_hub = mcp_cliente.Hub(J.VERSION, log=lambda m: J.log_terminal(m, "MCP"))
mcp_pending = {"tid": None, "args": None, "prompt": "", "expires": 0.0}


def mcp_trusted(tid):
    p = mcp_cliente.split_id(tid)
    trusted = {x.strip().lower() for x in J.settings["mcp_trusted"].split(",") if x.strip()}
    return bool(p) and p[0].lower() in trusted


def mcp_reconfigure():
    def work():
        try:
            J.mcp_hub.configure(J.settings["mcp_servers"])
        except ValueError as e:
            J.log_terminal(f"MCP: {e}", "AVISO")
        J.socketio.emit("mcp_status", J.mcp_hub.status())
    threading.Thread(target=work, daemon=True, name="mcp").start()


def run_mcp(tid, args, prompt="", confirmed=False):
    s = J.mcp_hub.find(tid)
    if not s:
        J.reply("Ese conector no está disponible ahora mismo.", "MCP")
        return
    server, tool = mcp_cliente.split_id(tid)
    if getattr(J._tl, "remote", False) and not J.mcp_trusted(tid):
        J.reply(f"Desde el celular solo uso los conectores marcados como de confianza, y «{server}» no lo está.", "MCP")
        return
    if not confirmed and not J.mcp_trusted(tid):
        J.mcp_pending.update(tid=tid, args=args, prompt=prompt, expires=time.time() + 30)
        J.reply(f"¿Uso «{tool}» del conector {server}? Di «confirmar» o «cancelar».", "MCP")
        return
    box = {}

    def work():
        try:
            box["out"] = J.mcp_hub.call(tid, args)
        except mcp_cliente.MCPError as e:
            box["err"] = str(e)
    th = threading.Thread(target=work, daemon=True)
    th.start()
    th.join(90)
    if th.is_alive():
        J.reply("El conector está tardando demasiado; lo dejo.", "MCP")
        return
    if "err" in box:
        J.reply(f"El conector {server} falló: {box['err']}", "MCP")
        return
    text, is_err = box["out"]
    diario.log("orden", f"conector {server}: {tool}")
    try:
        answer = J.ai_chat([{"role": "system", "content": J.settings["system_prompt"] + J.VOICE_RULE + J._user_ctx(prompt)},
                          {"role": "user", "content": f"Pedí: «{prompt}». La herramienta «{tool}» del conector {server} "
                                                      f"{'dio un error' if is_err else 'respondió'}:\n{text[:6000]}\n\n"
                                                      "Cuéntamelo en una a tres frases, como para decirlo en voz alta."}],
                         max_tokens=400).strip()
    except Exception:
        answer = text[:600] or "Hecho."
    J.reply(answer or "Hecho.", "MCP")


def mcp_answer(t):
    if not J.mcp_pending["tid"] or time.time() > J.mcp_pending["expires"]:
        return False
    tid, args, prompt = J.mcp_pending["tid"], J.mcp_pending["args"], J.mcp_pending["prompt"]
    if re.fullmatch(rf"{J.CONFIRM_RE}(?: .*)?", t):
        J.mcp_pending.update(tid=None, args=None, expires=0)
        J.run_mcp(tid, args, prompt, confirmed=True)
        return True
    if re.fullmatch(rf"{J.DENY_RE}(?: .*)?", t):
        J.mcp_pending.update(tid=None, args=None, expires=0)
        J.reply("Cancelado.", "MCP")
        return True
    return False


_MCP_LINK = os.path.join(J.CONFIG_DIR, "mcp_enlace.json")
mcp_server_token = {"value": ""}


def mcp_server_apply():
    """Jarvis como servidor MCP: escribe (o borra) el enlace que usa jarvis_mcp.py."""
    try:
        if J.settings["mcp_server_enabled"]:
            J.mcp_server_token["value"] = uuid.uuid4().hex + uuid.uuid4().hex
            os.makedirs(J.CONFIG_DIR, exist_ok=True)
            with open(J._MCP_LINK, "w", encoding="utf-8") as f:
                json.dump({"url": J.URL, "token": J.mcp_server_token["value"]}, f)
        else:
            J.mcp_server_token["value"] = ""
            if os.path.exists(J._MCP_LINK):
                os.remove(J._MCP_LINK)
    except OSError as e:
        J.log_terminal(f"Servidor MCP: {e}", "AVISO")


def jarvis_status_text():
    lines = [f"Jarvis {J.VERSION}. Micrófono: {J.state.get('mic')}. IA: {J.state.get('ai_label') or J.state.get('ai')}."]
    snap = J._timers_snapshot()
    if snap:
        lines.append("Temporizadores: " + "; ".join(f"{J._spoken_duration(x['remaining'])}" + (f" para {x['label']}" if x['label'] else "") for x in snap) + ".")
    try:
        items = J.agenda("hoy", with_day=False)
        if items:
            lines.append("Hoy: " + "; ".join(items[:6]) + ".")
    except Exception:
        pass
    last = [c for c in list(J.chat_buffer)[-6:]]
    if last:
        lines.append("Últimos mensajes: " + " | ".join(f"{c['role']}: {c['text'][:80]}" for c in last))
    return "\n".join(lines)


@J.app.route("/api/mcp", methods=["POST"])
def api_mcp():
    if request.remote_addr not in ("127.0.0.1", "::1"):
        abort(403)
    tok = request.headers.get("X-Jarvis-Token", "")
    if not (J.settings["mcp_server_enabled"] and J.mcp_server_token["value"]) or not hmac.compare_digest(tok, J.mcp_server_token["value"]):
        abort(403)
    d = request.get_json(silent=True) or {}
    tool, args = str(d.get("tool", "")), d.get("args") or {}
    texto = str(args.get("texto", "")).strip()[:1000]
    if tool == "jarvis_orden" and texto:
        J.log_terminal(f"Desde otra IA (MCP): {texto}", "MCP")
        return {"text": J.telegram_run(texto)}
    if tool == "jarvis_decir" and texto:
        J.reply(texto, "MCP")
        return {"text": "Dicho."}
    if tool == "jarvis_estado":
        return {"text": J.jarvis_status_text()}
    return {"text": "Herramienta desconocida o sin texto.", "error": True}


@J.socketio.on("mcp_info")
def handle_mcp_info(data=None):
    py = sys.executable.replace("pythonw.exe", "python.exe")
    cfg = {"mcpServers": {"jarvis": {"command": py, "args": [os.path.join(J.APP_DIR, "jarvis_mcp.py")]}}}
    return {"status": J.mcp_hub.status(), "claude_config": json.dumps(cfg, indent=2, ensure_ascii=False),
            "server": bool(J.settings["mcp_server_enabled"])}


@J.socketio.on("mcp_reload")
def handle_mcp_reload(data=None):
    J.mcp_reconfigure()
    return {"ok": True}


# =============================================================================
# 7.0 · PANEL EN EL CELULAR (misma wifi): página con código secreto y QR
# =============================================================================
mobile = {"server": None, "port": 0, "fails": {}, "last": 0.0}


def lan_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("10.255.255.255", 1))           # no envía nada: solo elige la interfaz de la red local
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        s.close()


def mobile_url():
    return f"http://{J.lan_ip()}:{J.settings['mobile_port']}/?k={J.settings['mobile_token']}"


mobile_app = Flask("jarvis_movil", template_folder=J.get_resource_path("templates"), static_folder=None)


def _mobile_ok():
    ip = request.remote_addr or "?"
    f = J.mobile["fails"].get(ip, (0, 0.0))
    if f[0] >= 20 and time.time() - f[1] < 600:
        abort(429)
    tok = request.headers.get("X-Jarvis-Token") or request.args.get("k", "")
    if not J.settings["mobile_token"] or not hmac.compare_digest(str(tok), J.settings["mobile_token"]):
        J.mobile["fails"][ip] = (f[0] + 1, time.time())
        abort(403)
    J.mobile["fails"].pop(ip, None)


@mobile_app.route("/")
def mobile_index():
    J._mobile_ok()
    return render_template("movil.html", version=J.VERSION)


@mobile_app.route("/api/estado")
def mobile_state():
    J._mobile_ok()
    try:
        items = J.agenda("hoy", with_day=False)
    except Exception:
        items = []
    return {"state": J.state, "timers": J._timers_snapshot(), "agenda": items[:8], "version": J.VERSION,
            "chat": list(J.chat_buffer)[-24:], "cards": list(J.recent_cards)[-4:], "name": J.settings["address_name"].strip()}


@mobile_app.route("/api/orden", methods=["POST"])
def mobile_order():
    J._mobile_ok()
    if time.time() - J.mobile["last"] < 1.0:
        return {"text": "Despacio: una orden cada vez."}
    J.mobile["last"] = time.time()
    texto = str((request.get_json(silent=True) or {}).get("texto", "")).strip()[:500]
    if not texto:
        return {"text": ""}
    J.push_chat("user", texto, "celular")
    J.log_terminal(f"Desde el panel del celular: {texto}", "CELULAR")
    out = J.telegram_run(texto)
    J.push_chat("jarvis", out, "celular")
    return {"text": out}


def mobile_apply():
    """Arranca, para o reinicia el panel del celular según los ajustes."""
    want = J.settings["mobile_enabled"] and bool(J.settings["mobile_token"])
    srv = J.mobile["server"]
    if srv and (not want or J.mobile["port"] != J.settings["mobile_port"]):
        try:
            srv.shutdown()
        except Exception:
            pass
        J.mobile.update(server=None, port=0)
    if want and not J.mobile["server"]:
        from werkzeug.serving import make_server
        try:
            srv = make_server("0.0.0.0", J.settings["mobile_port"], J.mobile_app, threaded=True)
        except OSError as e:
            J.log_terminal(f"Panel del celular: no pude abrir el puerto {J.settings['mobile_port']} ({e}).", "AVISO")
            return
        J.mobile.update(server=srv, port=J.settings["mobile_port"])
        threading.Thread(target=srv.serve_forever, daemon=True, name="panel-celular").start()
        J.log_terminal(f"Panel del celular en {J.mobile_url().split('?')[0]} (con tu código secreto).", "SISTEMA")


@J.socketio.on("mobile_info")
def handle_mobile_info(data=None):
    d = data or {}
    if d.get("new_token") or (J.settings["mobile_enabled"] and not J.settings["mobile_token"]):
        J.settings["mobile_token"] = uuid.uuid4().hex[:24]
        J.save_settings_file()
        J.mobile_apply()
    info = {"enabled": J.settings["mobile_enabled"], "running": bool(J.mobile["server"]),
            "url": J.mobile_url() if J.settings["mobile_token"] else ""}
    if info["url"]:
        try:
            import qrcode
            import io
            import base64
            img = qrcode.make(info["url"], border=2)
            buf = io.BytesIO()
            img.save(buf, "PNG")
            info["qr"] = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
        except Exception:
            info["qr"] = ""
    return info


# =============================================================================
# 7.0 · LEER EL TABLERO O UNA HOJA CON LA CÁMARA (y convertirlo en quiz o taller)
# =============================================================================
_BOARD_WHAT = r"(?:tablero|pizarra|pizarron|hoja|cuaderno|libro|pagina|lo que (?:te )?(?:muestro|ensen\w+)|esto que (?:te )?muestro)"
_BOARD_READ = re.compile(rf"\b(?:lee(?:me)?|transcribe|copia|pasa|anota|guarda|digitaliza|escanea)\b.*\b{_BOARD_WHAT}\b")
_BOARD_MAKE = re.compile(rf"\b(?:crea|haz|hazme|prepara|genera|convierte|arma)\b.*\b(?:quiz|evaluacion|examen|taller|preguntas|resumen)\b.*\b{_BOARD_WHAT}\b"
                         rf"|\bconvierte\b.*\b{_BOARD_WHAT}\b.*\b(?:quiz|evaluacion|examen|taller)\b")
_OCR_PROMPT = ("Transcribe fielmente TODO el texto escrito en esta imagen (un tablero, una hoja o un cuaderno), en el "
               "orden en que se lee. Conserva títulos, viñetas y fórmulas en texto plano (x², 3/4, √9). No añadas "
               "comentarios ni explicaciones. Si una parte no se lee, pon [ilegible].")


def board_voice(text, t):
    make = J._BOARD_MAKE.search(t)
    if not make and not J._BOARD_READ.search(t):
        return False
    if getattr(J._tl, "remote", False):
        J.reply("Leer el tablero o una hoja necesita que estés frente a la cámara del equipo.", "TABLERO")
        return True
    cfg, err = J._vision_cfg()
    if err:
        J.reply(err, "TABLERO")
        return True
    use_screen = bool(re.search(r"\bpantalla\b", t))
    J.reply("Muéstramelo a la cámara: en tres segundos tomo la foto." if not use_screen else "Leo lo que hay en la pantalla.", "TABLERO")

    def work():
        try:
            path = J.capture_screen_for_review() if use_screen else J.capture_camera(wait=3.0)
            transcript = "".join(J._vision_chunks(cfg, J._OCR_PROMPT, path)).strip()
            transcript = re.sub(r"<think>.*?</think>", "", transcript, flags=re.S).strip()
            if len(transcript) < 3:
                J.reply("No alcancé a leer nada. Acércalo un poco más a la cámara, con buena luz, y repite.", "TABLERO")
                return
            diario.log("orden", "leí el tablero/una hoja")
            if make:
                kind = aula.detect_kind(text) or "quiz"
                J.generate_material(f"{text}. Basado en este contenido (transcrito del tablero o la hoja):\n{transcript[:6000]}", kind=kind)
                return
            saved = ""
            if J.Document is not None:
                doc = J.Document()
                doc.add_heading(f"Tablero · {datetime.datetime.now():%d/%m/%Y %H:%M}", 1)
                try:
                    from docx.shared import Inches
                    doc.add_picture(path, width=Inches(6))
                except Exception:
                    pass
                for line in transcript.splitlines():
                    doc.add_paragraph(line)
                folder = os.path.join(J.settings["output_dir"], "Tablero")
                os.makedirs(folder, exist_ok=True)
                saved = os.path.join(folder, f"Tablero {datetime.datetime.now():%Y-%m-%d %H%M}.docx")
                doc.save(saved)
                J.generated_files.add(saved)
            if skills.pyperclip is not None:
                try:
                    skills.pyperclip.copy(transcript)
                except Exception:
                    pass
            J.card("tablero", "Lo que leí", transcript[:380], "▣", secs=20)
            J.reply("Listo: " + ("lo pasé a Word (carpeta Tablero) y " if saved else "") + "lo copié al portapapeles. "
                  + "Empieza así: " + re.sub(r"\s+", " ", transcript)[:160] + ("…" if len(transcript) > 160 else ""), "TABLERO")
            if saved and J.IS_WIN and J.settings["aula_open_docs"]:
                try:
                    os.startfile(saved)
                except OSError:
                    pass
        except ValueError as e:
            J.reply(str(e), "TABLERO")
        except Exception as e:
            J.log_terminal(f"Tablero: {e}", "ERROR")
            J.reply("No pude leerlo esta vez. Inténtalo de nuevo con más luz.", "TABLERO")
    threading.Thread(target=work, daemon=True, name="tablero").start()
    return True


__all__ = [
    '_BOARD_MAKE',
    '_BOARD_READ',
    '_BOARD_WHAT',
    '_CARD_ORIGINS',
    '_DIARY_RE',
    '_MCP_LINK',
    '_OCR_PROMPT',
    '_SPEED_RE',
    '_auto_card',
    '_downloads_dir',
    '_mobile_ok',
    '_partial_downloads',
    '_pro_cfg',
    '_pro_hourly',
    '_pro_say',
    '_pro_weather',
    '_proactive_worker',
    '_watch_worker',
    'api_mcp',
    'board_voice',
    'card',
    'diary_voice',
    'handle_habilidades_delete',
    'handle_habilidades_export',
    'handle_habilidades_folder',
    'handle_habilidades_get',
    'handle_habilidades_install',
    'handle_mcp_info',
    'handle_mcp_reload',
    'handle_mobile_info',
    'jarvis_status_text',
    'lan_ip',
    'mcp_answer',
    'mcp_hub',
    'mcp_pending',
    'mcp_reconfigure',
    'mcp_server_apply',
    'mcp_server_token',
    'mcp_trusted',
    'mobile',
    'mobile_app',
    'mobile_apply',
    'mobile_index',
    'mobile_order',
    'mobile_state',
    'mobile_url',
    'offer_routine_for',
    'pro_offer',
    'pro_state',
    'proactive_answer',
    'psutil',
    'recent_cards',
    'run_mcp',
    'skill_voice',
    'smart_greeting',
    'speed_voice',
    'watch_voice',
    'watches',
]
