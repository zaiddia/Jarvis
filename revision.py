"""Revisión de trabajos de estudiantes con la IA de visión.

Le muestras a Jarvis un taller, examen o cuaderno (cámara, pantalla o imagen) y devuelve:
aciertos, errores con su corrección, una nota sugerida (escala 1.0 a 5.0) y un mensaje de
retroalimentación para el estudiante. Se guarda un Word con la foto y un registro en CSV.

Este módulo no depende de la interfaz: jarvis.py lo conecta con la cámara, la voz y el tablero.
"""
import csv
import datetime
import os
import re


def build_prompt(que="", clave="", grado="", estudiante=""):
    p = ("Eres un docente colombiano experto revisando el trabajo escrito de un estudiante que "
         "aparece en esta imagen (foto de una hoja, cuaderno o captura de pantalla). Revisa con "
         "cuidado cada respuesta: comprueba los cálculos y el razonamiento, no solo el resultado. "
         "Sé justo y concreto. Escribe en español de Colombia, con notación matemática sencilla en "
         "texto plano (3/4, x², √9), sin LaTeX ni markdown.")
    if que:
        p += f" El docente dice que es: «{que}»."
    if grado:
        p += f" Grado: {grado}."
    if estudiante:
        p += f" Estudiante: {estudiante}."
    if clave:
        p += (" Usa estas respuestas correctas o criterios del docente como referencia principal: "
              f"«{clave}».")
    p += (" Si la imagen no se lee bien o no es un trabajo de estudiante, pon \"legible\": false y "
          "explica en \"resumen\" qué ves. La nota va de 1.0 a 5.0 (escala usada en Colombia; 3.0 "
          "aprueba) y es solo una sugerencia para el docente. La retroalimentación se dirige al "
          "estudiante: 2 o 3 frases, empezando por algo que hizo bien y luego qué mejorar. "
          "Responde SOLO con un objeto JSON con esta forma: "
          '{"legible": bool, "estudiante_en_hoja": str, "que_es": str, "resumen": str, '
          '"aciertos": [str], "errores": [{"donde": str, "que_esta_mal": str, "correccion": str}], '
          '"nota": number|null, "justificacion_nota": str, "retroalimentacion": str}')
    return p


def _txt(v, n=1500):
    return re.sub(r"\s+", " ", str(v or "")).strip()[:n]


def validate(d):
    if not isinstance(d, dict):
        raise ValueError("La IA no devolvió la revisión en el formato esperado.")
    nota = d.get("nota")
    try:
        nota = round(min(5.0, max(1.0, float(str(nota).replace(",", ".")))), 1) if nota not in (None, "") else None
    except (TypeError, ValueError):
        nota = None
    errores = []
    for e in d.get("errores") or []:
        if isinstance(e, dict) and (_txt(e.get("que_esta_mal")) or _txt(e.get("correccion"))):
            errores.append({"donde": _txt(e.get("donde"), 200), "que_esta_mal": _txt(e.get("que_esta_mal"), 600),
                            "correccion": _txt(e.get("correccion"), 600)})
        elif isinstance(e, str) and _txt(e):
            errores.append({"donde": "", "que_esta_mal": _txt(e, 600), "correccion": ""})
    legible = d.get("legible")
    legible = not (legible is False or str(legible).lower() == "false")
    return {"legible": legible, "estudiante_en_hoja": _txt(d.get("estudiante_en_hoja"), 80),
            "que_es": _txt(d.get("que_es"), 200), "resumen": _txt(d.get("resumen"), 800),
            "aciertos": [_txt(a, 400) for a in (d.get("aciertos") or []) if _txt(a)][:20],
            "errores": errores[:20], "nota": nota,
            "justificacion_nota": _txt(d.get("justificacion_nota"), 600),
            "retroalimentacion": _txt(d.get("retroalimentacion"), 900)}


def spoken_summary(r, estudiante=""):
    """Lo que Jarvis dice en voz alta (corto)."""
    if not r["legible"]:
        return ("No logro leer bien el trabajo. " + (r["resumen"] + " " if r["resumen"] else "") +
                "Acerca más la hoja a la cámara, con buena luz, o usa una foto.")
    quien = f" de {estudiante}" if estudiante else ""
    n_ok, n_err = len(r["aciertos"]), len(r["errores"])
    partes = [f"Revisé el trabajo{quien}: {n_ok} acierto{'s' if n_ok != 1 else ''} y "
              f"{n_err} error{'es' if n_err != 1 else ''}."]
    if r["nota"] is not None:
        partes.append(f"Nota sugerida: {str(r['nota']).replace('.', ',')}.")
    if r["errores"]:
        e = r["errores"][0]
        partes.append(f"Lo principal a corregir: {e['que_esta_mal']}")
    return " ".join(partes)


def safe_filename(text):
    t = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", str(text)).strip(" .")
    return re.sub(r"\s+", " ", t)[:60] or "trabajo"


def save_docx(r, image_path, out_dir, Document, estudiante="", curso="", que="", docente=""):
    from docx.shared import Pt, Cm, RGBColor
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    h = doc.add_paragraph()
    hr = h.add_run("Revisión de trabajo")
    hr.bold = True
    hr.font.size = Pt(16)
    info = [("Estudiante", estudiante or r["estudiante_en_hoja"] or "—"), ("Curso", curso or "—"),
            ("Trabajo", que or r["que_es"] or "—"), ("Fecha", datetime.date.today().strftime("%d/%m/%Y"))]
    t = doc.add_table(rows=1, cols=len(info))
    t.style = "Table Grid"
    for cell, (k, v) in zip(t.rows[0].cells, info):
        p = cell.paragraphs[0]
        p.add_run(f"{k}: ").bold = True
        p.add_run(v)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    if r["nota"] is not None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        run = p.add_run(f"Nota sugerida: {str(r['nota']).replace('.', ',')} / 5,0")
        run.bold = True
        run.font.size = Pt(14)
        if r["justificacion_nota"]:
            doc.add_paragraph(r["justificacion_nota"])
    if r["resumen"]:
        heading("Resumen")
        doc.add_paragraph(r["resumen"])
    if r["aciertos"]:
        heading("Aciertos")
        for a in r["aciertos"]:
            doc.add_paragraph(a, style="List Bullet")
    if r["errores"]:
        heading("Por corregir")
        tb = doc.add_table(rows=1, cols=3)
        tb.style = "Table Grid"
        for cell, txt in zip(tb.rows[0].cells, ["Dónde", "Qué está mal", "Corrección"]):
            cell.paragraphs[0].add_run(txt).bold = True
        for e in r["errores"]:
            row = tb.add_row().cells
            row[0].text, row[1].text, row[2].text = e["donde"], e["que_esta_mal"], e["correccion"]
    if r["retroalimentacion"]:
        heading("Mensaje para el estudiante")
        doc.add_paragraph(r["retroalimentacion"])
    if image_path and os.path.exists(image_path):
        heading("Trabajo revisado")
        try:
            doc.add_picture(image_path, width=Cm(16))
        except Exception:
            pass
    foot = doc.add_paragraph()
    fr = foot.add_run("Revisión sugerida por Jarvis con IA: la decisión final es del docente."
                      + (f" Docente: {docente}." if docente else ""))
    fr.italic = True
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    folder = os.path.join(out_dir, "Revisiones")
    os.makedirs(folder, exist_ok=True)
    stem = safe_filename(f"Revisión - {estudiante or r['estudiante_en_hoja'] or 'trabajo'} - {que or r['que_es'] or ''}")
    path = os.path.join(folder, f"{stem} {datetime.datetime.now():%Y-%m-%d %H%M}.docx")
    doc.save(path)
    _log_csv(folder, r, estudiante, curso, que, path)
    return path


def _log_csv(folder, r, estudiante, curso, que, path):
    """Registro acumulado (abre en Excel): una fila por trabajo revisado."""
    log = os.path.join(folder, "Registro de revisiones.csv")
    new = not os.path.exists(log)
    with open(log, "a", encoding="utf-8-sig" if new else "utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        if new:
            w.writerow(["Fecha", "Curso", "Estudiante", "Trabajo", "Nota sugerida", "Aciertos", "Errores", "Archivo"])
        w.writerow([datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), curso,
                    estudiante or r["estudiante_en_hoja"], que or r["que_es"],
                    str(r["nota"]).replace(".", ",") if r["nota"] is not None else "",
                    len(r["aciertos"]), len(r["errores"]), os.path.basename(path)])
