"""Valoraciones Periódicas de Desempeño como las calcula Zeti (10.1). Lógica pura, sin Jarvis.

La planilla de Zeti (zeti.co/calificaciones) tiene tres dimensiones y la autoevaluación, cada una con su peso:
Conocer 35 %, Hacer 30 %, Ser 30 % y Autoevaluación 5 %. En cada dimensión van varias actividades; su promedio
(DC, DH, DS, AEE) se multiplica por el peso (con dos decimales, redondeando hacia arriba desde 5) y la suma es la
valoración del periodo (Val, con un decimal) y su letra (VL: S, A, B, Bj). El acumulado del año (VF) pondera los
periodos (34 %, 33 %, 33 %) y VRA es lo que le falta al estudiante en lo que queda del año para llegar a 3,0.
"""
import re
import unicodedata
from decimal import ROUND_HALF_UP, Decimal

DIMS = [("conocer", "Conocer", "DC"), ("hacer", "Hacer", "DH"), ("ser", "Ser", "DS"), ("auto", "Autoevaluación", "AEE")]
WEIGHTS = {"conocer": 35, "hacer": 30, "ser": 30, "auto": 5}
SCALE = {"superior": 4.5, "alto": 4.0, "basico": 3.0}      # Superior 4,5–5,0 · Alto 4,0–4,4 · Básico 3,0–3,9 · Bajo 1,0–2,9
PASS = 3.0


def _plain(t):
    return "".join(c for c in unicodedata.normalize("NFD", str(t).lower()) if unicodedata.category(c) != "Mn")


def rhu(x, nd=1):
    """Redondeo como Zeti: 0,175 → 0,18 (no el redondeo «bancario» de Python)."""
    return float(Decimal(str(x)).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP))


_DIM_WORDS = [
    ("auto", r"auto\s?evaluaci"),
    ("ser", r"\b(?:ser|actitud|comportamiento|convivencia|participaci\w*|puntualidad|responsabilidad|disciplina|valores|"
            r"presentacion personal|respeto|compromiso|coevaluaci\w*)\b"),
    ("hacer", r"\b(?:hacer|taller\w*|practica\w*|laboratorio\w*|proyecto\w*|exposici\w*|trabajo\w*|tarea\w*|guia\w*|maqueta\w*|"
              r"ejercicio\w*|cuaderno|informe\w*|diseno|programaci\w*|actividad\w*|sustentaci\w*|investigaci\w*|cartelera\w*)\b"),
    ("conocer", r"\b(?:conocer|saber|quiz\w*|evaluaci\w*|examen\w*|prueba\w*|parcial\w*|test|previa\w*|leccion\w*|icfes|kahoot|"
                r"quizizz|forms|ticket|preguntas)\b"),
]


def infer_dim(actividad, dim=None):
    """La dimensión de una nota: la guardada o, si no, por el nombre de la actividad (por defecto, Conocer)."""
    if dim in WEIGHTS:
        return dim
    t = _plain(actividad)
    for d, rx in _DIM_WORDS:
        if re.search(rx, t):
            return d
    return "conocer"


def letter(val):
    if val is None:
        return ""
    return "S" if val >= SCALE["superior"] else "A" if val >= SCALE["alto"] else "B" if val >= SCALE["basico"] else "Bj"


def performance(val):
    return {"S": "Superior", "A": "Alto", "B": "Básico", "Bj": "Bajo", "": ""}[letter(val)]


def by_dim(notas):
    """{dim: [(actividad, nota)…]} en orden de fecha (varias notas de la misma actividad se promedian)."""
    out = {d: {} for d, _, _ in DIMS}
    for n in sorted(notas, key=lambda n: n.get("fecha", "")):
        if n.get("nota") is None:
            continue
        out[infer_dim(n.get("actividad", ""), n.get("dim"))].setdefault(n.get("actividad") or "nota", []).append(float(n["nota"]))
    return {d: [(a, sum(v) / len(v)) for a, v in acts.items()] for d, acts in out.items()}


def valoracion(notas, weights=None):
    """{"dims": {dim: promedio 1 decimal}, "partes": {dim: ponderado 2 decimales}, "val", "letra", "faltan": [dims sin notas]}."""
    w = weights or WEIGHTS
    bd = by_dim(notas)
    dims, partes, faltan = {}, {}, []
    for d, _, _ in DIMS:
        if not w.get(d):
            continue
        if bd[d]:
            dims[d] = rhu(sum(v for _, v in bd[d]) / len(bd[d]), 1)
            partes[d] = rhu(dims[d] * w[d] / 100, 2)
        else:
            faltan.append(d)
    if not partes:
        return {"dims": {}, "partes": {}, "val": None, "letra": "", "faltan": faltan}
    total_w = sum(w[d] for d in partes)
    s = sum(partes.values())
    val = rhu(s if total_w == 100 else s * 100 / total_w, 1)       # con dimensiones vacías, se reparte el peso
    return {"dims": dims, "partes": partes, "val": val, "letra": letter(val), "faltan": faltan}


def final(vals, period_weights=(34, 33, 33)):
    """VF: acumulado del año con las valoraciones de los periodos ya cerrados ([3.3, 4.2] → 2,5)."""
    s = sum(v * period_weights[i] / 100 for i, v in enumerate(vals) if v is not None and i < len(period_weights))
    return rhu(s, 1)


def vra(vf, done_periods, period_weights=(34, 33, 33)):
    """Lo que necesita en los periodos que faltan para llegar a 3,0 (None si ya aprobó)."""
    rest = sum(period_weights[done_periods:]) / 100
    if vf >= PASS or rest <= 0:
        return None
    return rhu((PASS - vf) / rest, 1)


def parse_weights(text, default=(35, 30, 30, 5)):
    nums = [int(x) for x in re.findall(r"\d{1,3}", str(text or ""))]
    if len(nums) != 4 or sum(nums) != 100:
        nums = list(default)
    return dict(zip([d for d, _, _ in DIMS], nums))


def parse_period_weights(text, default=(34, 33, 33)):
    nums = [int(x) for x in re.findall(r"\d{1,3}", str(text or ""))]
    return tuple(nums) if 2 <= len(nums) <= 6 and sum(nums) == 100 else default


# ----------------------------------------------------------------------------- planilla
def table(names, grades, since="", order="lista", sort_key=None, weights=None):
    """(actividades por dimensión, filas). Fila: {"nombre", "celdas": {dim: [nota|None…]}, "v": valoracion}."""
    since = since or ""
    acts = {d: [] for d, _, _ in DIMS}
    for n in names:
        for d, lst in by_dim([g for g in grades.get(n, []) if g.get("fecha", "") >= since]).items():
            for a, _ in lst:
                if a not in acts[d]:
                    acts[d].append(a)
    if order == "apellido" and sort_key:
        ordered = sorted(names, key=sort_key)
    elif order == "alfabetico":
        ordered = sorted(names, key=_plain)
    else:
        ordered = list(names)
    rows = []
    for n in ordered:
        gs = [g for g in grades.get(n, []) if g.get("fecha", "") >= since]
        bd = {d: dict(v) for d, v in by_dim(gs).items()}
        rows.append({"nombre": n, "celdas": {d: [rhu(bd[d][a], 1) if a in bd[d] else None for a in acts[d]] for d in acts},
                     "v": valoracion(gs, weights)})
    return acts, rows


def xlsx(curso, names, grades, absences, path, since="", order="lista", comma=True, institucion="", docente="", sort_key=None,
         weights=None, asignatura=""):
    """La planilla con la misma forma que «Valoraciones Periódicas de Desempeño» de Zeti."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    w = weights or WEIGHTS
    acts, rows = table(names, grades, since, order, sort_key, w)
    dims = [(d, name, code) for d, name, code in DIMS if w.get(d)]
    color = {"conocer": "D9F2D0", "hacer": "E4DCF7", "ser": "FBE7C6", "auto": "FFF2B3"}
    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "Valoraciones"
    ws["A1"] = "Valoraciones Periódicas de Desempeño · " + " · ".join(x for x in (curso, asignatura) if x)
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = " · ".join(x for x in (institucion, docente, f"notas desde {since}" if since else "todas las notas") if x)
    # fila 4: dimensiones (combinadas) · fila 5: actividades, promedio y porcentaje
    ws.cell(row=4, column=1, value="N°")
    ws.cell(row=4, column=2, value="Apellidos y Nombres de Estudiantes")
    ws.merge_cells(start_row=4, end_row=5, start_column=1, end_column=1)
    ws.merge_cells(start_row=4, end_row=5, start_column=2, end_column=2)
    col = 3
    spans = {}
    for d, name, code in dims:
        a = acts[d] if d != "auto" else []
        first = col
        for act in a:
            ws.cell(row=5, column=col, value=act[:28])
            col += 1
        ws.cell(row=5, column=col, value=code)
        ws.cell(row=5, column=col + 1, value=f"{w[d]}%")
        col += 2
        ws.cell(row=4, column=first, value=name)
        if col - 1 > first:
            ws.merge_cells(start_row=4, end_row=4, start_column=first, end_column=col - 1)
        spans[d] = (first, col - 1)
    tail = col
    for i, h in enumerate(("Val", "VL", "Fallas", "Observación")):
        ws.cell(row=4, column=tail + i, value=h)
        ws.merge_cells(start_row=4, end_row=5, start_column=tail + i, end_column=tail + i)
    for r in (4, 5):
        for c in range(1, tail + 4):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = box
    for d, (a, b) in spans.items():
        for r in (4, 5):
            for c in range(a, b + 1):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=color[d])
    for c in (1, 2):
        ws.cell(row=4, column=c).fill = PatternFill("solid", fgColor="D9F2D0")
    for i, row in enumerate(rows, 1):
        r = 5 + i
        ws.cell(row=r, column=1, value=f"{i:02d}")
        ws.cell(row=r, column=2, value=row["nombre"])
        v = row["v"]
        for d, _, _ in dims:
            a, b = spans[d]
            vals = row["celdas"][d] if d != "auto" else []
            for j, x in enumerate(vals):
                ws.cell(row=r, column=a + j, value=x).number_format = "0.0"
            ws.cell(row=r, column=b - 1, value=v["dims"].get(d)).number_format = "0.0"
            ws.cell(row=r, column=b, value=v["partes"].get(d)).number_format = "0.00"
        ws.cell(row=r, column=tail, value=v["val"]).number_format = "0.0"
        ws.cell(row=r, column=tail, value=v["val"]).font = Font(bold=True, color="C00000" if v["val"] is not None and v["val"] < PASS else "000000")
        ws.cell(row=r, column=tail + 1, value=v["letra"]).font = Font(bold=True)
        ws.cell(row=r, column=tail + 2, value=absences.get(row["nombre"], 0) or None)
        miss = [dict((x, y) for x, y, _ in DIMS)[d] for d in v["faltan"] if d != "auto" or acts["auto"]]
        if v["val"] is not None and miss:
            ws.cell(row=r, column=tail + 3, value="Sin notas de " + ", ".join(miss) + " (el peso se repartió)")
        for c in range(1, tail + 4):
            ws.cell(row=r, column=c).border = box
            if c > 2 and c < tail + 3:
                ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 38
    for c in range(3, tail + 3):
        ws.column_dimensions[get_column_letter(c)].width = 9
    ws.column_dimensions[get_column_letter(tail + 3)].width = 30
    for d, _, _ in dims:                               # las columnas de actividades, más anchas
        a, b = spans[d]
        for c in range(a, b - 1):
            ws.column_dimensions[get_column_letter(c)].width = 13
    ws.row_dimensions[5].height = 48
    ws.freeze_panes = ws.cell(row=6, column=3)
    ws.page_setup.orientation = "landscape"            # para imprimir: una hoja a lo ancho
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    # Para copiar: una columna por actividad, en el orden de la planilla de Zeti
    cp = wb.create_sheet("Para copiar")
    cp.append(["Copia cada columna (sin el título) y pégala en la misma actividad de Zeti, con la lista en el mismo orden."])
    heads, cols = ["Estudiante"], []
    for d, name, code in dims:
        if d == "auto":
            heads.append(f"{name} (AEE)")
            cols.append(("auto_dim", None))
            continue
        for j, act in enumerate(acts[d]):
            heads.append(f"{name}: {act[:24]}")
            cols.append((d, j))
    heads.append("Val")
    cp.append([])
    cp.append(heads)
    for c in cp[3]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")

    def fmt(x):
        if x is None:
            return ""
        s = f"{x:.1f}"
        return s.replace(".", ",") if comma else s
    for row in rows:
        line = [row["nombre"]]
        for d, j in cols:
            line.append(fmt(row["v"]["dims"].get("auto")) if d == "auto_dim" else fmt(row["celdas"][d][j]))
        line.append(fmt(row["v"]["val"]))
        cp.append(line)
    cp.column_dimensions["A"].width = 38
    cp.row_dimensions[3].height = 45
    wb.save(path)
    n_acts = sum(len(acts[d]) for d in acts)
    return len(rows), n_acts, rows
