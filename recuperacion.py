"""Planes de recuperación (8.8): actividades de nivelación para cada estudiante en riesgo y la carta al acudiente.

Con las notas bajas de cada estudiante (y sus fallas), la IA propone actividades concretas sobre lo que perdió, cómo
se valoran y una fecha de entrega, en el marco del Decreto 1290 (planes de apoyo para superar las debilidades). Todo
queda en Word para que lo revises y firmes.
"""
import datetime
import json
import os
import re


def weak(facts, umbral=3.0):
    """Las actividades que perdió (nota < umbral)."""
    return [n for n in facts["notas"] if n["nota"] < umbral]


def prompt(items, area="", grado="", dias=10):
    lines = [f"id {i}: " + json.dumps({"estudiante": f["estudiante"], "promedio": f["promedio"], "fallas": f["fallas"],
                                       "perdio": [f"{n['actividad']} ({n['nota']})" for n in weak(f)]}, ensure_ascii=False)
             for i, f in enumerate(items)]
    return ("Eres un docente de un colegio en Colombia que prepara planes de apoyo y recuperación (Decreto 1290). Para cada "
            "estudiante, con SOLO estos datos, propone 3 o 4 actividades de nivelación concretas y realizables en casa o en "
            "clase, centradas en los temas de las actividades que perdió (si el nombre de la actividad no dice el tema, "
            "propón repaso de lo trabajado en esa actividad). Para cada actividad di qué entrega el estudiante. Incluye "
            "cómo se valora (criterios claros, escala 1,0 a 5,0) y una recomendación para la familia. No inventes notas ni "
            f"diagnósticos. Plazo sugerido: {dias} días. " + (f"Área: {area}. " if area else "") + (f"Grado: {grado}. " if grado else "")
            + "\nDatos:\n" + "\n".join(lines)
            + '\n\nDevuelve SOLO JSON: {"planes": [{"id": 0, "temas": [str], "actividades": [{"titulo": str, "descripcion": str, '
              '"entrega": str}], "valoracion": str, "familia": str}]}')


def validate(d, n):
    out = {}
    for p in (d or {}).get("planes") or []:
        try:
            i = int(p["id"])
        except (KeyError, ValueError, TypeError):
            continue
        acts = [{"titulo": str(a.get("titulo", "")).strip()[:120], "descripcion": str(a.get("descripcion", "")).strip()[:900],
                 "entrega": str(a.get("entrega", "")).strip()[:300]}
                for a in (p.get("actividades") or []) if isinstance(a, dict) and str(a.get("titulo", "")).strip()]
        if 0 <= i < n and acts:
            out[i] = {"temas": [str(t)[:120] for t in (p.get("temas") or []) if str(t).strip()][:6], "actividades": acts[:6],
                      "valoracion": str(p.get("valoracion", "")).strip()[:900], "familia": str(p.get("familia", "")).strip()[:600]}
    return out


def _safe(text):
    return re.sub(r"\s+", " ", re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", str(text))).strip()[:70] or "recuperacion"


def build_docx(plans, curso, out_dir, Document, docente="", institucion="", dias=10):
    """plans: [(facts, plan)]. Un estudiante por página: plan de recuperación + carta y firmas."""
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    entrega = datetime.date.today() + datetime.timedelta(days=dias)
    for k, (f, p) in enumerate(plans):
        if k:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if institucion:
            doc.add_paragraph(institucion).runs[0].bold = True
        doc.add_heading(f"Plan de recuperación · {f['estudiante']}", 1)
        doc.add_paragraph(f"Curso: {curso}   ·   Promedio: {str(f['promedio']).replace('.', ',') if f['promedio'] is not None else '—'}"
                          f"   ·   Fallas: {f['fallas']}   ·   Entrega: {entrega:%d/%m/%Y}")
        perdidas = weak(f)
        if perdidas:
            doc.add_paragraph("Actividades por superar: " + "; ".join(f"{n['actividad']} ({str(n['nota']).replace('.', ',')})" for n in perdidas))
        if p["temas"]:
            doc.add_paragraph("Temas a reforzar: " + ", ".join(p["temas"]))
        tb = doc.add_table(rows=1, cols=3)
        tb.style = "Table Grid"
        for cell, h in zip(tb.rows[0].cells, ["Actividad", "Qué hacer", "Qué entregar"]):
            cell.paragraphs[0].add_run(h).bold = True
        for a in p["actividades"]:
            row = tb.add_row().cells
            row[0].text, row[1].text, row[2].text = a["titulo"], a["descripcion"], a["entrega"]
        if p["valoracion"]:
            doc.add_paragraph().add_run("Cómo se valora: ").bold = True
            doc.paragraphs[-1].add_run(p["valoracion"])
        doc.add_paragraph()
        doc.add_paragraph("Señor(a) acudiente:").runs[0].bold = True
        doc.add_paragraph(f"{f['estudiante']} tiene pendientes las actividades anteriores. Le pedimos acompañar este plan "
                          f"y verificar que lo entregue a más tardar el {entrega:%d/%m/%Y}. " + (p["familia"] or ""))
        doc.add_paragraph("\nFirma del docente: ______________________   " + (docente or "") +
                          "\nFirma del estudiante: ____________________\nFirma del acudiente: _____________________")
    folder = os.path.join(out_dir, "Planes de recuperación")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{_safe(curso if len(plans) > 1 else plans[0][0]['estudiante'])} {datetime.date.today():%Y-%m-%d}.docx")
    doc.save(path)
    return path
