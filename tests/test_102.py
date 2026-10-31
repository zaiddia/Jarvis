"""Pruebas de la 10.2: acumulado del año (VF) y lo que falta para aprobar (VRA), como en Zeti."""
import os

import pytest
from openpyxl import load_workbook

import periodos
import seguimiento

PEGADO = """Apellidos y Nombres de Estudiantes\tNov\t1715\t\t\tDC\t35%
01 Agredo Gueguia Adrian\tNvo\t3.5\t\t\t3.5\t1.23\t3.3\t\t\t3.3\t0.99\t3.0\t\t\t3.0\t0.90\t3.5\t0.18\t3.3\tB\t2.5\t1.5
04 Camacho Viera Jheferson David\t\t4.8\t\t\t4.8\t1.68\t4.6\t\t\t4.6\t1.38\t4.6\t\t\t4.6\t1.38\t4.5\t0.23\t4.7\tS\t3.1\t
05 Dagua Dagua Arley Josue\tRei\t\t\t\t\t\t\t\t\t\t1.0\t3.0
10 Muñoz Malte Angela Maria 4.6 4.6 1.61 4.4 4.4 1.32 4.5 4.5 1.35 4.0 0.20 4.5 S 2.9 0.3"""


def test_pegar_la_tabla_de_zeti():
    rows = periodos.parse_paste(PEGADO)
    assert rows == [{"nombre": "Agredo Gueguia Adrian", "val": 3.3, "vf": 2.5},
                    {"nombre": "Camacho Viera Jheferson David", "val": 4.7, "vf": 3.1},
                    {"nombre": "Muñoz Malte Angela Maria", "val": 4.5, "vf": 2.9}]


@pytest.fixture
def anio(jv, monkeypatch, tmp_path):
    monkeypatch.setattr(periodos, "_path", str(tmp_path / "periodos.json"))
    monkeypatch.setattr(periodos, "data", {})
    monkeypatch.setattr(seguimiento, "_path", str(tmp_path / "notas.json"))
    monkeypatch.setattr(seguimiento, "data", {})
    monkeypatch.setitem(jv.aula.data, "cursos", {})
    monkeypatch.setitem(jv.aula.data, "asistencia", {})
    monkeypatch.setattr(jv.aula, "save", lambda: None)
    monkeypatch.setattr(jv, "save_settings_file", lambda: None)
    monkeypatch.setitem(jv.settings, "period_current", 3)
    monkeypatch.setitem(jv.settings, "zeti_period_weights", "34,33,33")
    monkeypatch.setitem(jv.settings, "zeti_since", "2026-09-01")
    monkeypatch.setitem(jv.settings, "output_dir", str(tmp_path / "out"))
    return tmp_path


def _notas(c, h, s, a="2026-10-01"):
    return [{"actividad": "Quiz", "nota": c, "fecha": a}, {"actividad": "Taller", "nota": h, "fecha": a}, {"actividad": "Actitud", "nota": s, "fecha": a}]


def test_acumulado_y_quien_puede_perder(jv, anio, monkeypatch):
    r = jv.handle_periods_paste({"curso": "10A", "periodo": 1, "text": PEGADO})
    assert r["ok"] and r["n"] == 3 and r["creado"]                                # crea la lista con la de Zeti
    assert jv.aula.data["cursos"]["10A"][0] == "Agredo Gueguia Adrian"
    for name, v in (("Agredo Gueguia Adrian", 4.2), ("Camacho Viera Jheferson David", 4.5), ("Muñoz Malte Angela Maria", 4.0)):
        periodos.set_val("10A", name, 2, v)
    seguimiento.data["10A"] = {"Agredo Gueguia Adrian": _notas(1.0, 1.0, 1.5),
                               "Camacho Viera Jheferson David": _notas(4.0, 4.0, 4.0),
                               "Muñoz Malte Angela Maria": _notas(2.0, 2.0, 2.0)}
    s = jv.year_status("10A", "Agredo Gueguia Adrian")
    assert s["vf"] == 2.5 and s["vra"] == 1.5 and s["actual"] == 1.2 and s["pierde"]       # igual que en Zeti
    assert jv.year_status("10A", "Camacho Viera Jheferson David")["aprobo"]
    m = jv.year_status("10A", "Muñoz Malte Angela Maria")
    assert m["vf"] == 2.9 and m["vra"] == 0.3 and not m["pierde"]
    jv._rec.clear()
    jv.process_command("¿quiénes pueden perder el año en 10A?")
    assert "pueden perder el año 1: Agredo Gueguia Adrian (necesita 1,5, va en 1,2)" in jv._rec.said()
    jv._rec.clear()
    jv.process_command("¿cuánto necesita Agredo de 10A para pasar?")
    assert "lleva 2,5 acumulado y necesita 1,5" in jv._rec.said() and "no le alcanza" in jv._rec.said()
    jv._rec.clear()
    jv.process_command("¿quiénes ya ganaron el año en 10A?")
    assert "Camacho" in jv._rec.said()
    assert any("pueden perder el año en 10A" in x["t"] for x in jv.attention_items())
    # la planilla para Zeti trae VF y VRA
    monkeypatch.setattr(jv, "_open_doc", lambda p: False)
    jv.process_command("exporta las notas de 10A para Zeti")
    path = [os.path.join(r, f) for r, _, fs in os.walk(jv.settings["output_dir"]) for f in fs if f.startswith("Zeti")][0]
    ws = load_workbook(path)["Valoraciones"]
    head = [c.value for c in ws[4]]
    i = head.index("VF")
    assert head[i + 1] == "VRA"
    row = [c.value for c in ws[6]]
    assert row[1] == "Agredo Gueguia Adrian" and row[i] == 2.5 and row[i + 1] == 1.5
    assert [c.value for c in ws[7]][i + 1] == "✓"


def test_falta_un_periodo_y_cambiar_de_periodo(jv, anio):
    periodos.set_val("7A", "Juan Pérez", 1, 3.0)
    jv.aula.data["cursos"]["7A"] = ["Juan Pérez"]
    jv._rec.clear()
    jv.process_command("¿cuánto necesita Juan de 7A para pasar?")
    assert "periodo 2" in jv._rec.said()
    jv.process_command("estamos en el segundo periodo")
    assert jv.settings["period_current"] == 2
    s = jv.year_status("7A", "Juan Pérez")
    assert s["vf"] == 1.0 and s["vra"] == 3.0                                       # (3,0 - 1,0) / 0,66


def test_al_cerrar_el_periodo_guarda_la_val(jv, anio):
    jv.aula.data["cursos"]["7A"] = ["Ana"]
    seguimiento.data["7A"] = {"Ana": _notas(4.0, 4.0, 4.0)}
    jv.settings["period_current"] = 1
    assert jv.save_period_vals() == 1
    assert periodos.data["7A"]["Ana"] == {"1": 4.0} and jv.settings["period_current"] == 2


def test_pesos_de_periodos(jv):
    assert jv._validate("zeti_period_weights", "25 25 25 25") == "25,25,25,25"
    with pytest.raises(ValueError):
        jv._validate("zeti_period_weights", "34,33")
