"""Pruebas de la 8.7: Google Calendar y Drive directos, y exportar a Kahoot, Quizizz y Zeti."""

import pytest
from openpyxl import load_workbook

import calendario
import classroom
import exportar
import gservicios
import seguimiento

QUIZ = {"titulo": "Evaluación: fracciones", "tema": "fracciones",
        "preguntas": [{"enunciado": "¿Cuánto es 1/2 + 1/2?" + " muy largo" * 20, "opciones": ["1", "2", "1/4", "0"],
                       "respuesta": "a", "explicacion": "Suma de mitades."},
                      {"enunciado": "¿Mitad de 8?", "opciones": ["2", "4", "6", "8"], "respuesta": "b", "explicacion": ""}]}


class FakeGoogle:
    def __init__(self):
        self.calls = []

    def call(self, method, path, params=None, body=None, base=classroom.API):
        self.calls.append((method, base, path, params, body))
        if base == gservicios.CAL_API and method == "POST":
            return {"id": "e1", "conferenceData": {"entryPoints": [{"entryPointType": "video", "uri": "https://meet.google.com/abc"}]}}
        if base == gservicios.CAL_API:
            return {"items": [{"id": "x", "summary": "Consejo académico", "start": {"dateTime": "2026-11-03T14:00:00-05:00"},
                               "end": {"dateTime": "2026-11-03T15:00:00-05:00"}},
                              {"id": "y", "summary": "Festivo", "start": {"date": "2026-11-02"}, "end": {"date": "2026-11-03"}}]}
        if path == "files" and method == "POST":
            return {"id": "FOLDER"}
        if path.endswith("/permissions"):
            return {}
        return {"id": "F", "webViewLink": "https://drive.google.com/file/d/F", "trashed": False}

    def raw(self, method, url, data, ctype):
        self.calls.append((method, url, ctype, data[:400]))
        return {"id": "F2", "name": "Quiz", "webViewLink": "https://docs.google.com/presentation/d/F2"}


def _google(jv, monkeypatch):
    g = FakeGoogle()
    monkeypatch.setattr(jv, "classroom_client", lambda: g)
    monkeypatch.setitem(jv.settings, "google_scopes", " ".join(classroom.SCOPES))
    return g


def test_cuerpo_del_evento():
    ev = {"titulo": "Reunión con padres", "fecha": "2026-11-05", "hora": "15:00", "duracion_min": 45, "lugar": "Salón 3", "descripcion": ""}
    b = gservicios.event_body(ev, meet=True)
    assert b["start"]["dateTime"] == "2026-11-05T15:00:00" and b["end"]["dateTime"] == "2026-11-05T15:45:00"
    assert b["start"]["timeZone"] == "America/Bogota" and "createRequest" in b["conferenceData"]
    b = gservicios.event_body(dict(ev, hora=""))
    assert b["start"] == {"date": "2026-11-05"} and b["end"] == {"date": "2026-11-06"}


def test_evento_directo_en_google_calendar(jv, monkeypatch):
    g = _google(jv, monkeypatch)
    jv._rec.clear()
    ev = {"titulo": "Reunión con padres de Juan", "fecha": "2026-11-05", "hora": "15:00", "duracion_min": 60, "lugar": "",
          "descripcion": "", "meet": True}
    jv.event_create(ev)
    assert "Google Calendar" in jv._rec.said() and "meet.google.com/abc" in jv._rec.said()
    assert any("conferenceDataVersion=1" in c[2] for c in g.calls)


def test_lee_tu_google_calendar(jv, monkeypatch):
    _google(jv, monkeypatch)
    monkeypatch.setitem(jv.settings, "calendar_urls", "")
    monkeypatch.setattr(calendario, "cache", {"events": [], "fetched": 0.0, "errors": [], "urls": ()})
    jv._gcal["t"] = 0.0
    assert jv.cal_configured()
    jv.cal_refresh(force=True)
    titles = [e["titulo"] for e in calendario.cache["events"]]
    assert titles == ["Festivo", "Consejo académico"] and calendario.cache["events"][0]["todo_el_dia"]


def test_subir_a_drive_y_compartir(jv, monkeypatch, tmp_path):
    g = _google(jv, monkeypatch)
    saved = {}
    monkeypatch.setattr(jv, "apply_settings", lambda d: saved.update(d) or jv.settings.update(d))
    monkeypatch.setitem(jv.settings, "drive_folder_id", "")
    monkeypatch.setitem(jv.settings, "drive_auto", False)
    monkeypatch.setattr(jv.webbrowser if hasattr(jv, "webbrowser") else __import__("webbrowser"), "open", lambda u: None)
    import nucleo.gintegra as gi
    monkeypatch.setattr(gi.webbrowser, "open", lambda u: None)
    p = tmp_path / "Presentación fracciones.pptx"
    p.write_bytes(b"PK..")
    jv.generated_files.add(str(p))
    said = []
    monkeypatch.setattr(jv, "notify", lambda text, origin="JARVIS": said.append(text))
    jv._tl.capture = []
    try:
        jv.process_command("pásalo a Google Slides")
    finally:
        jv._tl.capture = None
    up = [c for c in g.calls if c[0] == "POST" and str(c[1]).startswith(gservicios.UPLOAD)][-1]
    assert b"application/vnd.google-apps.presentation" in up[3] and b'"parents": ["FOLDER"]' in up[3]
    assert saved["drive_folder_id"] == "FOLDER" and "Drive" in said[-1]
    jv._tl.capture = []
    try:
        jv.process_command("súbelo a Drive y compártelo")
    finally:
        jv._tl.capture = None
    assert jv.aprobaciones.latest()["titulo"].startswith("Compartir")
    for a in jv.aprobaciones.pending():
        jv.resolve_approval(a["id"], False, "prueba")


def test_kahoot_y_quizizz(tmp_path):
    rows = exportar.kahoot_rows(QUIZ, seconds=25)
    assert len(rows[0][0]) <= 120 and rows[0][5] == 20 and rows[0][6] == "1" and rows[1][6] == "2"
    wb = load_workbook(exportar.kahoot_xlsx(QUIZ, str(tmp_path)))
    ws = wb.active
    assert ws["B8"].value.startswith("Question") and ws["A9"].value == 1 and ws["G9"].value == 20 and ws["H10"].value == "2"
    wb = load_workbook(exportar.quizizz_xlsx(QUIZ, str(tmp_path)))
    ws = wb.active
    assert [c.value for c in ws[1]][:3] == ["Question Text", "Question Type", "Option 1"]
    assert ws["B2"].value == "Multiple Choice" and ws["H2"].value == 1 and ws["K2"].value == "Suma de mitades."


def test_kahoot_por_voz(jv, monkeypatch, tmp_path):
    monkeypatch.setitem(jv.settings, "output_dir", str(tmp_path))
    monkeypatch.setattr(jv, "_open_doc", lambda p: False)
    jv.remember_quiz(QUIZ)
    jv._rec.clear()
    jv.process_command("pasa el quiz a Kahoot")
    assert "Kahoot" in jv._rec.said() and list((tmp_path / "Exportaciones").glob("Kahoot*.xlsx"))


def test_planilla_para_zeti(jv, monkeypatch, tmp_path):
    monkeypatch.setitem(jv.settings, "output_dir", str(tmp_path))
    monkeypatch.setitem(jv.settings, "zeti_since", "2026-08-01")
    monkeypatch.setitem(jv.settings, "zeti_order", "apellido")
    monkeypatch.setitem(jv.settings, "zeti_comma", True)
    monkeypatch.setattr(jv, "_open_doc", lambda p: False)
    monkeypatch.setitem(jv.aula.data, "cursos", {"7A": ["Juan Pérez", "Ana Gómez", "Luis Rojas"]})
    monkeypatch.setitem(jv.aula.data, "activo", "7A")
    monkeypatch.setitem(jv.aula.data, "asistencia", {"7A": {"2026-09-01": ["Juan Pérez"]}})
    monkeypatch.setattr(seguimiento, "data", {"7A": {
        "Juan Pérez": [{"actividad": "Quiz 1", "nota": 2.0, "fecha": "2026-09-01"}, {"actividad": "Taller", "nota": 3.0, "fecha": "2026-09-10"},
                       {"actividad": "Viejo", "nota": 5.0, "fecha": "2026-05-01"}],
        "Ana Gómez": [{"actividad": "Quiz 1", "nota": 4.6, "fecha": "2026-09-01"}]}})
    jv._rec.clear()
    jv.process_command("exporta las notas de 7A para Zeti")
    assert "3 estudiantes y 2 actividades" in jv._rec.said() and "Conocer (35 %)" in jv._rec.said()
    path = next((tmp_path / "Exportaciones").glob("Zeti 7A*.xlsx"))
    wb = load_workbook(path)
    ws = wb["Valoraciones"]
    assert [c.value for c in ws[4]][:3] == ["N°", "Apellidos y Nombres de Estudiantes", "Conocer"]
    assert [c.value for c in ws[5]][2:12] == ["Quiz 1", "DC", "35%", "Taller", "DH", "30%", "DS", "30%", "AEE", "5%"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=6)]
    assert [r[1] for r in rows] == ["Ana Gómez", "Juan Pérez", "Luis Rojas"]                  # por apellido
    assert rows[1][2:8] == [2.0, 2.0, 0.7, 3.0, 3.0, 0.9] and rows[1][12:15] == [2.5, "Bj", 1]
    assert rows[0][12:14] == [4.6, "S"] and "Sin notas de Hacer, Ser" in rows[0][15]
    cp = [[c.value for c in r] for r in wb["Para copiar"].iter_rows(min_row=3)]
    assert cp[0] == ["Estudiante", "Conocer: Quiz 1", "Hacer: Taller", "Autoevaluación (AEE)", "Val"]
    assert cp[2][:3] == ["Juan Pérez", "2,0", "3,0"] and cp[2][4] == "2,5"


def test_zeti_desde_valida_la_fecha(jv):
    assert jv._validate("zeti_since", "2026-08-03") == "2026-08-03"
    with pytest.raises(ValueError):
        jv._validate("zeti_since", "el lunes")


