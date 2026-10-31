"""Herramientas de aula para Jarvis.

· Cursos: listas de estudiantes guardadas en aula.json (carpeta de configuración).
· Sorteo justo: nadie repite hasta que hayan pasado todos; se salta a quien faltó hoy.
· Grupos aleatorios equilibrados (de N personas o N grupos).
· Asistencia del día, exportable a CSV.
· Generador de material (quiz, taller, rúbrica, plan de clase) a Word, con la IA.

Este módulo no depende de la interfaz: jarvis.py lo conecta con la voz y el tablero.
"""
import csv
import datetime
import difflib
import json

import datos
import math
import os
import random
import re
import threading
import unicodedata

_lock = threading.Lock()
_path = None
data = {"cursos": {}, "activo": "", "pendientes": {}, "asistencia": {}}

ORDINALES = {"primero": "1", "primer": "1", "segundo": "2", "tercero": "3", "tercer": "3",
             "cuarto": "4", "quinto": "5", "sexto": "6", "septimo": "7", "octavo": "8",
             "noveno": "9", "decimo": "10", "once": "11", "undecimo": "11", "onceavo": "11",
             "uno": "1", "dos": "2", "tres": "3", "cuatro": "4", "cinco": "5", "seis": "6",
             "siete": "7", "ocho": "8", "nueve": "9", "diez": "10"}
NUM_WORDS = {"un": 1, "una": 1, "uno": 1, "dos": 2, "tres": 3, "cuatro": 4, "cinco": 5, "seis": 6,
             "siete": 7, "ocho": 8, "nueve": 9, "diez": 10, "once": 11, "doce": 12, "quince": 15,
             "veinte": 20, "treinta": 30}


def normalize(text):
    t = unicodedata.normalize("NFD", str(text).lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^\w\s]", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def to_int(word):
    word = normalize(word)
    if word.isdigit():
        return int(word)
    return NUM_WORDS.get(word)


# ---------------------------------------------------------------------------
# Guardado
# ---------------------------------------------------------------------------
def configure(config_dir):
    global _path
    _path = os.path.join(config_dir, "aula.json")
    load()


locked = False      # 8.9: el archivo está cifrado con otra clave: no se lee ni se sobrescribe


def load():
    global locked
    try:
        d = datos.read_json(_path)
        locked = False
        if isinstance(d, dict):
            for k in data:
                if isinstance(d.get(k), type(data[k])):
                    data[k] = d[k]
    except datos.DataLocked:
        locked = True
    except (ValueError, TypeError, OSError):
        pass


def save():
    if not _path or locked:
        return
    datos.write_json(_path, data)


def snapshot():
    with _lock:
        cur = data["activo"]
        return {"cursos": {k: list(v) for k, v in data["cursos"].items()}, "activo": cur,
                "ausentes": list(_ausentes(cur)), "pendientes": len(data["pendientes"].get(cur, [])),
                "hoy": _today()}


# ---------------------------------------------------------------------------
# Cursos
# ---------------------------------------------------------------------------
def parse_names(text):
    """Una persona por línea. Acepta lo pegado desde Excel (toma la 1ª columna con texto)
    y quita numeraciones como «1.», «2)» o «03 -»."""
    out, seen = [], set()
    for line in str(text or "").splitlines():
        cells = [c.strip() for c in line.split("\t")]
        cell = next((c for c in cells if re.search(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]", c)), "")
        name = re.sub(r"^\s*\d+\s*[.)\-–:]?\s*", "", cell).strip(" ,;")
        name = re.sub(r"\s+", " ", name)
        if name and normalize(name) not in seen:
            seen.add(normalize(name))
            out.append(name)
    return out


def course_key(name):
    """«Séptimo A», «7 a», «7A» -> «7a» (para reconocer el curso por voz)."""
    words = [ORDINALES.get(w, w) for w in normalize(name).replace("grado", " ").replace("curso", " ").split()]
    return "".join(words)


def save_course(name, names_text, old_name=None):
    name = str(name or "").strip()[:40]
    if not name:
        raise ValueError("Ponle un nombre al curso, por ejemplo 7A.")
    names = parse_names(names_text)
    with _lock:
        if old_name and old_name != name and old_name in data["cursos"]:
            data["cursos"].pop(old_name)
            data["pendientes"].pop(old_name, None)
            data["asistencia"][name] = data["asistencia"].pop(old_name, {})
            if data["activo"] == old_name:
                data["activo"] = name
        data["cursos"][name] = names
        data["pendientes"].pop(name, None)            # la lista cambió: el sorteo empieza de nuevo
        if not data["activo"]:
            data["activo"] = name
        save()
    return names


def delete_course(name):
    with _lock:
        data["cursos"].pop(name, None)
        data["pendientes"].pop(name, None)
        data["asistencia"].pop(name, None)
        if data["activo"] == name:
            data["activo"] = next(iter(data["cursos"]), "")
        save()


def find_course(spoken):
    key = course_key(spoken)
    if not key:
        return None
    keys = {course_key(c): c for c in data["cursos"]}
    if key in keys:
        return keys[key]
    m = difflib.get_close_matches(key, list(keys), n=1, cutoff=0.75)
    return keys[m[0]] if m else None


def set_active(name):
    with _lock:
        if name not in data["cursos"]:
            raise ValueError(f"No tengo el curso «{name}».")
        data["activo"] = name
        save()


def active():
    """(nombre del curso, lista) o lanza ValueError con un mensaje para decir en voz alta."""
    cur = data["activo"]
    if not data["cursos"]:
        raise ValueError("Aún no tengo ninguna lista de estudiantes. Ábrela en el tablero del aula, "
                         "pestaña Cursos, y pega los nombres.")
    if cur not in data["cursos"]:
        cur = data["activo"] = next(iter(data["cursos"]))
    if not data["cursos"][cur]:
        raise ValueError(f"El curso {cur} no tiene estudiantes todavía.")
    return cur, data["cursos"][cur]


# ---------------------------------------------------------------------------
# Asistencia del día
# ---------------------------------------------------------------------------
def _today():
    return datetime.date.today().isoformat()


def _ausentes(course):
    return set(data["asistencia"].get(course, {}).get(_today(), []))


def find_student(spoken, names):
    """Busca a un estudiante por cómo se dijo su nombre. Devuelve (nombre, candidatos)."""
    q = normalize(spoken)
    if not q:
        return None, []
    exact = [n for n in names if normalize(n) == q]
    if exact:
        return exact[0], exact
    starts = [n for n in names if normalize(n).startswith(q + " ") or q in normalize(n).split()]
    if len(starts) == 1:
        return starts[0], starts
    if len(starts) > 1:
        full = [n for n in starts if q in normalize(n)]
        return None, full or starts
    scored = sorted(((difflib.SequenceMatcher(None, q, normalize(n)).ratio(), n) for n in names), reverse=True)
    firsts = sorted(((difflib.SequenceMatcher(None, q, normalize(n).split()[0]).ratio(), n) for n in names),
                    reverse=True)
    best = max(scored[:1] + firsts[:1]) if names else (0, None)
    return (best[1], [best[1]]) if best[0] >= 0.78 else (None, [])


def mark(name, present):
    with _lock:
        cur, names = active()
        if name not in names:
            raise ValueError(f"{name} no está en {cur}.")
        day = data["asistencia"].setdefault(cur, {}).setdefault(_today(), [])
        if present and name in day:
            day.remove(name)
        elif not present and name not in day:
            day.append(name)
        # Solo se guardan los últimos 120 días de cada curso
        hist = data["asistencia"][cur]
        for old in sorted(hist)[:-120]:
            hist.pop(old)
        save()
        return cur, len(names) - len(day), len(names)


def export_attendance(out_dir):
    cur, names = active()
    aus = _ausentes(cur)
    folder = os.path.join(out_dir, "Asistencia")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"Asistencia {safe_filename(cur)} {_today()}.csv")
    # utf-8-sig: Excel abre bien las tildes; «;» es el separador de Excel en español
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Curso", "Fecha", "Estudiante", "Estado"])
        for n in names:
            w.writerow([cur, _today(), n, "Ausente" if n in aus else "Presente"])
    return path, len(names) - len(aus), len(names)


# ---------------------------------------------------------------------------
# Sorteo justo y grupos
# ---------------------------------------------------------------------------
def pick():
    """Escoge a alguien que aún no haya pasado (y que no haya faltado hoy).
    Devuelve (nombre, quedan, total_presentes)."""
    with _lock:
        cur, names = active()
        present = [n for n in names if n not in _ausentes(cur)]
        if not present:
            raise ValueError("Hoy están todos marcados como ausentes.")
        pool = [n for n in data["pendientes"].get(cur, []) if n in present]
        if not pool:                                   # ya pasaron todos: nueva ronda
            pool = present[:]
            random.shuffle(pool)
        name = pool.pop()
        data["pendientes"][cur] = pool
        save()
        return name, len(pool), len(present)


def reset_pick():
    with _lock:
        data["pendientes"].pop(data["activo"], None)
        save()


def make_groups(size=None, count=None):
    """Grupos aleatorios equilibrados (tamaños que difieren como mucho en 1)."""
    with _lock:
        cur, names = active()
        present = [n for n in names if n not in _ausentes(cur)]
    if len(present) < 2:
        raise ValueError("Necesito al menos dos estudiantes presentes para hacer grupos.")
    if count:
        n_groups = max(1, min(int(count), len(present)))
    else:
        size = max(2, int(size or 4))
        n_groups = max(1, round(len(present) / size))
        # que ningún grupo quede mucho más grande de lo pedido
        while n_groups < len(present) and math.ceil(len(present) / n_groups) > size + 1:
            n_groups += 1
    people = present[:]
    random.shuffle(people)
    groups = [[] for _ in range(n_groups)]
    for i, p in enumerate(people):
        groups[i % n_groups].append(p)
    return cur, groups


# ---------------------------------------------------------------------------
# Generador de material de clase (Word)
# ---------------------------------------------------------------------------
KINDS = {
    "quiz": ("Evaluación", r"quiz|evaluacion|examen|prueba|cuestionario|test"),
    "taller": ("Taller", r"taller|guia|ficha|ejercicios|actividad|tarea"),
    "rubrica": ("Rúbrica", r"rubrica"),
    "plan": ("Plan de clase", r"plan de clase|planeacion|planeador|planificacion|secuencia didactica"),
    # 8.2
    "saber": ("Preguntas tipo Saber", r"(?:preguntas|prueba|simulacro|quiz|evaluacion)\s+(?:tipo\s+)?(?:saber|icfes)(?:\s+\d+)?|"
                                      r"tipo saber|tipo icfes|simulacro icfes|simulacro saber"),
    "malla": ("Malla curricular", r"malla(?: curricular)?|plan (?:del|de|para el) periodo|planeacion (?:del|de) periodo|"
                                  r"plan de area|plan de asignatura|planeacion anual|plan anual"),
    "piar": ("Ajustes razonables (PIAR)", r"piar|ajustes razonables|adaptacion curricular|flexibilizacion curricular|"
                                          r"para (?:un|una|el|la) (?:estudiante|alumno|alumna|nino|nina) (?:con|que tiene)"),
}
DEFAULT_COUNT = {"quiz": 10, "taller": 8, "rubrica": 4, "plan": 0, "saber": 8, "malla": 10, "piar": 0}
NIVELES = ["Superior", "Alto", "Básico", "Bajo"]      # escala de valoración en Colombia


def detect_kind(text):
    t = normalize(text)
    for kind in ("piar", "saber", "malla", "rubrica", "plan", "quiz", "taller"):
        if re.search(rf"\b(?:{KINDS[kind][1]})\b", t):
            return kind
    return None


def detect_count(text):
    t = normalize(text)
    m = re.search(r"\b(\d+|" + "|".join(NUM_WORDS) + r")\s+(?:preguntas|ejercicios|puntos|problemas|criterios|items)\b", t)
    return to_int(m.group(1)) if m else None


def build_prompt(kind, request, count=None, grade="", topic=""):
    n = max(1, min(40, count or DEFAULT_COUNT[kind] or 1))
    base = ("Eres un docente experto en Colombia que prepara material de clase claro, correcto y "
            "adecuado a la edad. Escribe en español de Colombia. Usa notación matemática sencilla en "
            "texto plano (3/4, x², √9, 2·5), sin LaTeX ni markdown. ")
    pedido = f"Petición del docente: «{request.strip()}»."
    if topic:
        pedido += f" Tema: {topic}."
    if grade:
        pedido += f" Grado: {grade}."
    if kind == "quiz":
        schema = ('{"titulo": str, "grado": str, "tema": str, "instrucciones": str, '
                  '"preguntas": [{"enunciado": str, "opciones": [str, str, str, str], '
                  '"respuesta": "a"|"b"|"c"|"d", "explicacion": str}]}')
        task = (f"Crea una evaluación de {n} preguntas de opción múltiple con 4 opciones cada una y una "
                "sola correcta. Varía la posición de la respuesta correcta y el nivel de dificultad. "
                "Las opciones van sin la letra delante. La explicación es una frase para el docente.")
    elif kind == "taller":
        schema = ('{"titulo": str, "grado": str, "tema": str, "objetivo": str, '
                  '"para_recordar": str, "ejemplo": str, '
                  '"ejercicios": [{"enunciado": str, "respuesta": str}]}')
        task = (f"Crea un taller con {n} ejercicios graduados de menor a mayor dificultad. "
                "«para_recordar» es una explicación breve del tema (3 a 5 frases) y «ejemplo» un "
                "ejemplo resuelto paso a paso. «respuesta» es la solución para la clave del docente.")
    elif kind == "rubrica":
        schema = ('{"titulo": str, "grado": str, "tema": str, '
                  '"criterios": [{"criterio": str, "descriptores": [str, str, str, str]}]}')
        task = (f"Crea una rúbrica de evaluación con {n} criterios. Los 4 descriptores de cada criterio "
                "corresponden, en este orden, a los niveles de desempeño Superior, Alto, Básico y Bajo "
                "(escala de valoración usada en Colombia). Descriptores concretos y observables.")
    elif kind == "saber":
        schema = ('{"titulo": str, "grado": str, "tema": str, "instrucciones": str, '
                  '"preguntas": [{"contexto": str, "enunciado": str, "opciones": [str, str, str, str], '
                  '"respuesta": "a"|"b"|"c"|"d", "competencia": str, "explicacion": str}]}')
        task = (f"Crea {n} preguntas tipo Saber (estilo ICFES) de selección múltiple con única respuesta. Cada pregunta "
                "parte de un «contexto» propio (una situación, un texto corto, una tabla o un gráfico descrito en texto) y "
                "evalúa una competencia del área tal como la nombra el ICFES (por ejemplo, en matemáticas: interpretación y "
                "representación, formulación y ejecución, argumentación). Las opciones incorrectas deben ser errores "
                "frecuentes y creíbles. La explicación dice por qué la correcta lo es. No copies preguntas reales de "
                "pruebas publicadas.")
    elif kind == "malla":
        schema = ('{"titulo": str, "grado": str, "tema": str, "periodo": str, "estandar": str, '
                  '"semanas": [{"semana": int, "tema": str, "dba": str, "desempeno": str, "actividades": str, '
                  '"evaluacion": str}], "evidencias": [str]}')
        task = (f"Crea la malla curricular (planeación del periodo) de {n} semanas. «tema» del objeto es el área o "
                "asignatura. Para cada semana: el tema, el DBA o estándar básico de competencias del MEN que se trabaja (con "
                "número solo si lo conoces con certeza; si no, descríbelo sin inventar el número), un desempeño observable, "
                "actividades sugeridas y cómo se evalúa. Ordena los temas de lo básico a lo complejo. «evidencias» son las "
                "evidencias de aprendizaje del periodo.")
    elif kind == "piar":
        schema = ('{"titulo": str, "grado": str, "tema": str, "necesidad": str, "barreras": [str], '
                  '"ajustes": [{"tipo": str, "ajuste": str}], "actividad": str, "evaluacion": str, "seguimiento": str}')
        task = ("Prepara ajustes razonables (en el marco del PIAR, Decreto 1421 de 2017) para que un estudiante con la "
                "situación que describe el docente participe en la clase o el material indicado. No diagnostiques ni uses "
                "lenguaje clínico: habla de barreras para el aprendizaje y la participación y de apoyos. «necesidad» resume "
                "en lenguaje respetuoso lo que dijo el docente. «ajustes» incluye ajustes de acceso, de la actividad, de "
                "tiempos, de materiales y de evaluación («tipo» es uno de esos). «actividad» es la versión adaptada y "
                "concreta de la actividad o el material, lista para usar. «evaluacion» dice cómo valorar el avance con "
                "esos ajustes y «seguimiento» qué observar y cuándo revisar con la familia y orientación.")
    else:
        schema = ('{"titulo": str, "grado": str, "tema": str, "duracion": str, "objetivo": str, '
                  '"competencias": [str], "momentos": [{"nombre": str, "duracion": str, '
                  '"actividades": str}], "recursos": [str], "evaluacion": str}')
        task = ("Crea un plan de clase con los momentos inicio (exploración), desarrollo "
                "(estructuración y práctica) y cierre (transferencia y valoración), con tiempos "
                "realistas que sumen la duración de la clase (55 minutos si no se indica otra). "
                "Alinéalo con los Derechos Básicos de Aprendizaje (DBA) y los Estándares Básicos de Competencias "
                "del MEN de Colombia para ese grado y área: en «competencias» nombra el DBA o estándar que se "
                "trabaja (con su número si lo conoces con certeza; si no, descríbelo sin inventar el número).")
    return (base + pedido + " " + task + " Responde SOLO con un objeto JSON con esta forma: " + schema)


def _txt(v, limit=2000):
    return re.sub(r"\s+", " ", str(v or "")).strip()[:limit]


def validate(kind, d, request=""):
    """Limpia y comprueba lo que devolvió la IA. Lanza ValueError si no sirve."""
    if not isinstance(d, dict):
        raise ValueError("La IA no devolvió el material en el formato esperado.")
    out = {"titulo": _txt(d.get("titulo"), 140) or KINDS[kind][0],
           "grado": _txt(d.get("grado"), 40), "tema": _txt(d.get("tema"), 140)}
    if kind in ("quiz", "saber"):
        qs = []
        for q in d.get("preguntas") or []:
            if not isinstance(q, dict):
                continue
            opts = [re.sub(r"^\s*[a-dA-D][).:]\s*", "", _txt(o, 400)) for o in (q.get("opciones") or [])][:4]
            ans = normalize(q.get("respuesta"))[:1]
            if _txt(q.get("enunciado")) and len(opts) == 4 and all(opts) and ans in "abcd" and ans:
                qs.append({"enunciado": _txt(q["enunciado"], 800), "opciones": opts, "respuesta": ans,
                           "explicacion": _txt(q.get("explicacion"), 400)})
                if kind == "saber":
                    qs[-1].update(contexto=_txt(q.get("contexto"), 2500), competencia=_txt(q.get("competencia"), 200))
        if not qs:
            raise ValueError("La IA no generó preguntas válidas.")
        out.update(instrucciones=_txt(d.get("instrucciones"), 500) or
                   "Lee cada pregunta y marca con una X la única respuesta correcta.", preguntas=qs)
    elif kind == "taller":
        ej = [{"enunciado": _txt(e.get("enunciado"), 800), "respuesta": _txt(e.get("respuesta"), 800)}
              for e in (d.get("ejercicios") or []) if isinstance(e, dict) and _txt(e.get("enunciado"))]
        if not ej:
            raise ValueError("La IA no generó ejercicios válidos.")
        out.update(objetivo=_txt(d.get("objetivo"), 500), para_recordar=_txt(d.get("para_recordar"), 2500),
                   ejemplo=_txt(d.get("ejemplo"), 2000), ejercicios=ej)
    elif kind == "rubrica":
        crit = []
        for c in d.get("criterios") or []:
            if isinstance(c, dict) and _txt(c.get("criterio")):
                desc = [_txt(x, 500) for x in (c.get("descriptores") or [])][:4]
                desc += [""] * (4 - len(desc))
                crit.append({"criterio": _txt(c["criterio"], 200), "descriptores": desc})
        if not crit:
            raise ValueError("La IA no generó criterios válidos.")
        out.update(criterios=crit)
    elif kind == "malla":
        sem = []
        for i, s in enumerate(d.get("semanas") or [], 1):
            if isinstance(s, dict) and _txt(s.get("tema")):
                try:
                    num = int(s.get("semana") or i)
                except (TypeError, ValueError):
                    num = i
                sem.append({"semana": num, "tema": _txt(s["tema"], 300), "dba": _txt(s.get("dba"), 600),
                            "desempeno": _txt(s.get("desempeno"), 600), "actividades": _txt(s.get("actividades"), 1200),
                            "evaluacion": _txt(s.get("evaluacion"), 600)})
        if not sem:
            raise ValueError("La IA no generó las semanas del periodo.")
        out.update(periodo=_txt(d.get("periodo"), 60), estandar=_txt(d.get("estandar"), 800), semanas=sem,
                   evidencias=[_txt(x, 400) for x in (d.get("evidencias") or []) if _txt(x)])
    elif kind == "piar":
        aj = [{"tipo": _txt(a.get("tipo"), 60) or "Ajuste", "ajuste": _txt(a.get("ajuste"), 800)}
              for a in (d.get("ajustes") or []) if isinstance(a, dict) and _txt(a.get("ajuste"))]
        if not aj:
            raise ValueError("La IA no generó ajustes válidos.")
        out.update(necesidad=_txt(d.get("necesidad"), 800), barreras=[_txt(b, 400) for b in (d.get("barreras") or []) if _txt(b)],
                   ajustes=aj, actividad=_txt(d.get("actividad"), 4000), evaluacion=_txt(d.get("evaluacion"), 1500),
                   seguimiento=_txt(d.get("seguimiento"), 1200))
    else:
        mom = [{"nombre": _txt(m.get("nombre"), 80), "duracion": _txt(m.get("duracion"), 40),
                "actividades": _txt(m.get("actividades"), 2500)}
               for m in (d.get("momentos") or []) if isinstance(m, dict) and _txt(m.get("nombre"))]
        if not mom:
            raise ValueError("La IA no generó los momentos de la clase.")
        out.update(duracion=_txt(d.get("duracion"), 40), objetivo=_txt(d.get("objetivo"), 800),
                   competencias=[_txt(c, 300) for c in (d.get("competencias") or []) if _txt(c)],
                   momentos=mom, recursos=[_txt(r, 200) for r in (d.get("recursos") or []) if _txt(r)],
                   evaluacion=_txt(d.get("evaluacion"), 1500))
    return out


def safe_filename(text):
    t = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", str(text)).strip(" .")
    return re.sub(r"\s+", " ", t)[:70] or "material"


def build_docx(kind, m, out_dir, Document, docente="", institucion="", curso=""):
    """Escribe el .docx y devuelve su ruta."""
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt, Cm, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    if kind in ("rubrica", "malla"):                       # la rúbrica y la malla caben mejor en horizontal
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width

    # Encabezado institucional
    hdr = [("Institución", institucion), ("Docente", docente),
           ("Grado", m.get("grado") or curso), ("Fecha", "")]
    hdr = [(k, v) for k, v in hdr if v or k == "Fecha"]
    t = doc.add_table(rows=1, cols=len(hdr))
    t.style = "Table Grid"
    for cell, (k, v) in zip(t.rows[0].cells, hdr):
        p = cell.paragraphs[0]
        r = p.add_run(f"{k}: ")
        r.bold = True
        p.add_run(v or "____________")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(14)
    tr = title.add_run(m["titulo"])
    tr.bold = True
    tr.font.size = Pt(16)
    if m.get("tema"):
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(f"Tema: {m['tema']}")
        sr.italic = True
    if kind in ("quiz", "taller", "saber"):
        doc.add_paragraph("Nombre del estudiante: _______________________________________   "
                          "Curso: ________")

    def heading(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.bold = True
        r.font.size = Pt(12.5)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        return h

    def key_page(title_text):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        heading(title_text)

    if kind == "quiz":
        heading("Instrucciones")
        doc.add_paragraph(m["instrucciones"])
        for i, q in enumerate(m["preguntas"], 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            p.add_run(f"{i}. ").bold = True
            p.add_run(q["enunciado"])
            for letter, opt in zip("abcd", q["opciones"]):
                op = doc.add_paragraph(f"{letter}) {opt}")
                op.paragraph_format.left_indent = Cm(0.8)
                op.paragraph_format.space_after = Pt(0)
        key_page("Clave de respuestas (solo para el docente)")
        for i, q in enumerate(m["preguntas"], 1):
            p = doc.add_paragraph()
            p.add_run(f"{i}. {q['respuesta']})").bold = True
            if q["explicacion"]:
                p.add_run(f"  {q['explicacion']}")
    elif kind == "saber":
        heading("Instrucciones")
        doc.add_paragraph(m["instrucciones"])
        for i, q in enumerate(m["preguntas"], 1):
            if q.get("contexto"):
                c = doc.add_paragraph(q["contexto"])
                c.paragraph_format.space_before = Pt(10)
                c.paragraph_format.keep_with_next = True
                c.runs[0].italic = True
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run(f"{i}. ").bold = True
            p.add_run(q["enunciado"])
            for letter, opt in zip("ABCD", q["opciones"]):
                op = doc.add_paragraph(f"{letter}. {opt}")
                op.paragraph_format.left_indent = Cm(0.8)
                op.paragraph_format.space_after = Pt(0)
        key_page("Clave, competencias y justificación (solo para el docente)")
        for i, q in enumerate(m["preguntas"], 1):
            p = doc.add_paragraph()
            p.add_run(f"{i}. {q['respuesta'].upper()}").bold = True
            if q.get("competencia"):
                p.add_run(f"  ·  {q['competencia']}").italic = True
            if q["explicacion"]:
                p.add_run(f"  {q['explicacion']}")
    elif kind == "malla":
        info = [x for x in (f"Grado: {m['grado']}" if m["grado"] else "", f"Periodo: {m['periodo']}" if m.get("periodo") else "") if x]
        if info:
            doc.add_paragraph("   ·   ".join(info))
        if m.get("estandar"):
            heading("Estándar básico de competencias")
            doc.add_paragraph(m["estandar"])
        tb = doc.add_table(rows=1, cols=6)
        tb.style = "Table Grid"
        for cell, txt in zip(tb.rows[0].cells, ["Semana", "Tema", "DBA / estándar", "Desempeño", "Actividades", "Evaluación"]):
            cell.paragraphs[0].add_run(txt).bold = True
        for s in m["semanas"]:
            row = tb.add_row().cells
            for cell, v in zip(row, [str(s["semana"]), s["tema"], s["dba"], s["desempeno"], s["actividades"], s["evaluacion"]]):
                cell.text = v
        if m.get("evidencias"):
            heading("Evidencias de aprendizaje del periodo")
            for e in m["evidencias"]:
                doc.add_paragraph(e, style="List Bullet")
    elif kind == "piar":
        if m.get("necesidad"):
            heading("Situación del estudiante")
            doc.add_paragraph(m["necesidad"])
        if m.get("barreras"):
            heading("Barreras para el aprendizaje y la participación")
            for b in m["barreras"]:
                doc.add_paragraph(b, style="List Bullet")
        heading("Ajustes razonables")
        tb = doc.add_table(rows=1, cols=2)
        tb.style = "Table Grid"
        for cell, txt in zip(tb.rows[0].cells, ["Tipo", "Ajuste"]):
            cell.paragraphs[0].add_run(txt).bold = True
        for a in m["ajustes"]:
            row = tb.add_row().cells
            row[0].text, row[1].text = a["tipo"], a["ajuste"]
        if m.get("actividad"):
            heading("Actividad adaptada")
            doc.add_paragraph(m["actividad"])
        if m.get("evaluacion"):
            heading("Valoración con ajustes")
            doc.add_paragraph(m["evaluacion"])
        if m.get("seguimiento"):
            heading("Seguimiento")
            doc.add_paragraph(m["seguimiento"])
        doc.add_paragraph("Ajustes en el marco del Decreto 1421 de 2017. Revísalos con orientación escolar y la familia.").runs[0].italic = True
    elif kind == "taller":
        if m["objetivo"]:
            heading("Objetivo")
            doc.add_paragraph(m["objetivo"])
        if m["para_recordar"]:
            heading("Para recordar")
            doc.add_paragraph(m["para_recordar"])
        if m["ejemplo"]:
            heading("Ejemplo resuelto")
            doc.add_paragraph(m["ejemplo"])
        heading("Ejercicios")
        for i, e in enumerate(m["ejercicios"], 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.add_run(f"{i}. ").bold = True
            p.add_run(e["enunciado"])
            for _ in range(3):                               # espacio para responder
                doc.add_paragraph("_" * 80).paragraph_format.space_after = Pt(0)
        key_page("Soluciones (solo para el docente)")
        for i, e in enumerate(m["ejercicios"], 1):
            p = doc.add_paragraph()
            p.add_run(f"{i}. ").bold = True
            p.add_run(e["respuesta"] or "—")
    elif kind == "rubrica":
        tb = doc.add_table(rows=1, cols=5)
        tb.style = "Table Grid"
        for cell, txt in zip(tb.rows[0].cells, ["Criterio", *NIVELES]):
            cell.paragraphs[0].add_run(txt).bold = True
        for c in m["criterios"]:
            row = tb.add_row().cells
            row[0].paragraphs[0].add_run(c["criterio"]).bold = True
            for cell, d in zip(row[1:], c["descriptores"]):
                cell.text = d
        doc.add_paragraph()
        doc.add_paragraph("Escala de valoración: Superior, Alto, Básico y Bajo.").runs[0].italic = True
    else:
        info = [x for x in (f"Grado: {m['grado']}" if m["grado"] else "",
                            f"Duración: {m['duracion']}" if m["duracion"] else "") if x]
        if info:
            doc.add_paragraph("   ·   ".join(info))
        heading("Objetivo de aprendizaje")
        doc.add_paragraph(m["objetivo"] or "—")
        if m["competencias"]:
            heading("Competencias")
            for c in m["competencias"]:
                doc.add_paragraph(c, style="List Bullet")
        heading("Momentos de la clase")
        tb = doc.add_table(rows=1, cols=3)
        tb.style = "Table Grid"
        for cell, txt in zip(tb.rows[0].cells, ["Momento", "Tiempo", "Actividades"]):
            cell.paragraphs[0].add_run(txt).bold = True
        for mo in m["momentos"]:
            row = tb.add_row().cells
            row[0].text, row[1].text, row[2].text = mo["nombre"], mo["duracion"], mo["actividades"]
        if m["recursos"]:
            heading("Recursos")
            for r in m["recursos"]:
                doc.add_paragraph(r, style="List Bullet")
        if m["evaluacion"]:
            heading("Evaluación")
            doc.add_paragraph(m["evaluacion"])

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = foot.add_run(f"Preparado con Jarvis · revísalo antes de usarlo · {datetime.date.today():%d/%m/%Y}")
    fr.italic = True
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    folder = os.path.join(out_dir, "Material de clase")
    os.makedirs(folder, exist_ok=True)
    stem = safe_filename(f"{KINDS[kind][0]} - {m.get('tema') or m['titulo']}")
    path = os.path.join(folder, f"{stem} {datetime.datetime.now():%Y-%m-%d %H%M}.docx")
    doc.save(path)
    return path
