"""Informes para padres y observador del estudiante (8.2).

El observador guarda anotaciones por estudiante (positivas, académicas y de convivencia) con fecha. El informe para
padres junta lo que Jarvis sabe del estudiante (notas del seguimiento, fallas de la asistencia y el observador) y la
IA redacta una carta respetuosa con fortalezas, aspectos por mejorar y recomendaciones para la casa. Nada se envía
solo: quedan en Word para que los revises.
"""
import datetime
import json
import os
import re
import threading
import unicodedata

import datos

_lock = threading.Lock()
_path = None
data = {}           # {curso: {estudiante: [{"fecha", "tipo", "texto"}]}}


def norm(t):
    t = unicodedata.normalize("NFD", str(t).lower())
    return re.sub(r"\s+", " ", "".join(c for c in t if unicodedata.category(c) != "Mn")).strip()


def configure(config_dir):
    global _path
    _path = os.path.join(config_dir, "observador.json")
    load()


locked = False      # 8.9: cifrado con otra clave: no se lee ni se sobrescribe


def load():
    global data, locked
    try:
        d = datos.read_json(_path)
        data, locked = (d if isinstance(d, dict) else {}), False
    except datos.DataLocked:
        data, locked = {}, True
    except (OSError, ValueError, TypeError):
        data = {}


def save():
    if not _path or locked:
        return
    try:
        datos.write_json(_path, data)
    except OSError:
        pass


def kind_of(text):
    t = norm(text)
    if re.search(r"\b(?:felicit|destac|excelente|muy bien|ayudo|lidera|mejoro|gran|sobresal|colabor)", t):
        return "positiva"
    if re.search(r"\b(?:conviv|pelea|peleo|irrespet|agred|agresi|grit|insult|celular|uniforme|llego tarde|evadi|fuga|bullying|matoneo|grosero)", t):
        return "convivencia"
    return "académica"


def add(curso, estudiante, texto, fecha=None, tipo=None):
    texto = str(texto).strip()
    texto = texto[:1].upper() + texto[1:]
    with _lock:
        data.setdefault(curso, {}).setdefault(estudiante, []).append(
            {"fecha": (fecha or datetime.date.today()).isoformat(), "tipo": tipo or kind_of(texto), "texto": texto[:1000]})
        save()
    return data[curso][estudiante][-1]


def entries(curso, estudiante):
    return list(data.get(curso, {}).get(estudiante, []))


def spoken(curso, estudiante):
    e = entries(curso, estudiante)
    if not e:
        return f"{estudiante} no tiene anotaciones en el observador de {curso}."
    by = {}
    for x in e:
        by[x["tipo"]] = by.get(x["tipo"], 0) + 1
    last = e[-1]
    return (f"{estudiante} tiene {len(e)} {'anotaciones' if len(e) != 1 else 'anotación'} en {curso} ("
            + ", ".join(f"{n} {k}{'s' if n != 1 else ''}" for k, n in by.items())
            + f"). La última, del {last['fecha']}: {last['texto']}")


def performance(avg):
    """Desempeño en la escala nacional (Decreto 1290): rangos habituales de las instituciones."""
    if avg is None:
        return ""
    return "Superior" if avg >= 4.6 else "Alto" if avg >= 4.0 else "Básico" if avg >= 3.0 else "Bajo"


def facts(curso, estudiante, notas, fallas):
    avg = round(sum(n["nota"] for n in notas) / len(notas), 2) if notas else None
    return {"estudiante": estudiante, "curso": curso, "promedio": avg, "desempeno": performance(avg), "fallas": fallas,
            "notas": [{"actividad": n["actividad"], "nota": n["nota"], "fecha": n.get("fecha", "")} for n in notas[-15:]],
            "observador": [{"fecha": x["fecha"], "tipo": x["tipo"], "texto": x["texto"]} for x in entries(curso, estudiante)[-10:]]}


def letter_prompt(items, docente="", institucion=""):
    """Prompt para redactar el informe de uno o varios estudiantes (lista de facts)."""
    lines = [f"id {i}: " + json.dumps(f, ensure_ascii=False) for i, f in enumerate(items)]
    return ("Eres un docente de un colegio en Colombia que escribe a los padres o acudientes. Para cada estudiante, con "
            "SOLO estos datos (no inventes notas, hechos ni diagnósticos), redacta un informe breve, cálido y respetuoso: "
            "empieza por las fortalezas, explica con claridad lo que debe mejorar (citando notas o fallas si las hay) y da "
            "2 o 3 recomendaciones concretas para apoyar desde la casa. Si no hay datos suficientes, dilo con tacto y "
            "propón una reunión. Trato de «usted». "
            + (f"Docente: {docente}. " if docente else "") + (f"Institución: {institucion}. " if institucion else "")
            + "\nDatos:\n" + "\n".join(lines)
            + '\n\nDevuelve SOLO JSON: {"informes": [{"id": 0, "parrafos": [str, str], "recomendaciones": [str]}]}')


def letter_validate(d, n):
    out = {}
    for r in (d or {}).get("informes") or []:
        try:
            i = int(r["id"])
        except (KeyError, ValueError, TypeError):
            continue
        par = [str(p).strip()[:1500] for p in (r.get("parrafos") or []) if str(p).strip()]
        if 0 <= i < n and par:
            out[i] = {"parrafos": par[:5], "recomendaciones": [str(x).strip()[:400] for x in (r.get("recomendaciones") or []) if str(x).strip()][:5]}
    return out


def _safe(text):
    return re.sub(r"\s+", " ", re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", str(text))).strip()[:70] or "informe"


def _base_doc(Document):
    from docx.shared import Pt, Cm
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    return doc


def build_letters(reports, out_dir, Document, docente="", institucion="", title=""):
    """reports: [(facts, texto)]. Un estudiante por página. Devuelve la ruta."""
    from docx.enum.text import WD_BREAK
    doc = _base_doc(Document)
    for k, (f, txt) in enumerate(reports):
        if k:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if institucion:
            doc.add_paragraph(institucion).runs[0].bold = True
        doc.add_paragraph(f"{datetime.date.today():%d/%m/%Y}")
        doc.add_heading(f"Informe para la familia de {f['estudiante']}", 1)
        doc.add_paragraph(f"Curso: {f['curso']}" + (f"   ·   Promedio: {str(f['promedio']).replace('.', ',')} ({f['desempeno']})"
                                                    if f["promedio"] is not None else "") + f"   ·   Fallas: {f['fallas']}")
        if f["notas"]:
            tb = doc.add_table(rows=1, cols=3)
            tb.style = "Table Grid"
            for cell, h in zip(tb.rows[0].cells, ["Actividad", "Nota", "Fecha"]):
                cell.paragraphs[0].add_run(h).bold = True
            for n in f["notas"]:
                row = tb.add_row().cells
                row[0].text, row[1].text, row[2].text = n["actividad"], str(n["nota"]).replace(".", ","), n["fecha"]
            doc.add_paragraph()
        doc.add_paragraph("Estimada familia:")
        for p in txt["parrafos"]:
            doc.add_paragraph(p)
        if txt["recomendaciones"]:
            doc.add_paragraph("Recomendaciones para la casa:").runs[0].bold = True
            for r in txt["recomendaciones"]:
                doc.add_paragraph(r, style="List Bullet")
        doc.add_paragraph("\nCordialmente,\n\n_____________________________\n" + (docente or "Docente"))
    folder = os.path.join(out_dir, "Informes para padres")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{_safe(title or reports[0][0]['estudiante'])} {datetime.date.today():%Y-%m-%d}.docx")
    doc.save(path)
    return path


def build_observer(curso, estudiante, out_dir, Document, institucion=""):
    doc = _base_doc(Document)
    if institucion:
        doc.add_paragraph(institucion).runs[0].bold = True
    doc.add_heading(f"Observador del estudiante: {estudiante}", 1)
    doc.add_paragraph(f"Curso: {curso}   ·   Generado el {datetime.date.today():%d/%m/%Y}")
    tb = doc.add_table(rows=1, cols=3)
    tb.style = "Table Grid"
    for cell, h in zip(tb.rows[0].cells, ["Fecha", "Tipo", "Anotación"]):
        cell.paragraphs[0].add_run(h).bold = True
    for x in entries(curso, estudiante):
        row = tb.add_row().cells
        row[0].text, row[1].text, row[2].text = x["fecha"], x["tipo"].capitalize(), x["texto"]
    doc.add_paragraph("\nFirma del docente: ____________________    Firma del estudiante: ____________________\n"
                      "Firma del acudiente: ____________________")
    folder = os.path.join(out_dir, "Observador")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{_safe(estudiante + ' ' + curso)} {datetime.date.today():%Y-%m-%d}.docx")
    doc.save(path)
    return path
