"""Presentaciones (PowerPoint) y hojas de Excel por voz.

    «haz una presentación de 8 diapositivas sobre el ciclo del agua para quinto»
        -> la IA escribe el contenido (JSON) y aquí se arma el .pptx, con notas para el docente.
    «crea la planilla de notas de 7A con 4 notas»
        -> .xlsx con la lista del curso, fórmulas de promedio y desempeño (escala editable).
    «crea una tabla en Excel de los planetas con su distancia al sol»
        -> la IA arma la tabla y aquí se da formato.

jarvis.py decide cuándo usarlo; este módulo no habla ni llama a la IA.
"""
import datetime
import os
import re
import unicodedata

NUM_WORDS = {"un": 1, "una": 1, "uno": 1, "dos": 2, "tres": 3, "cuatro": 4, "cinco": 5, "seis": 6,
             "siete": 7, "ocho": 8, "nueve": 9, "diez": 10, "once": 11, "doce": 12, "trece": 13,
             "catorce": 14, "quince": 15, "veinte": 20}
GRADOS = {"primero": "1°", "segundo": "2°", "tercero": "3°", "cuarto": "4°", "quinto": "5°", "sexto": "6°",
          "septimo": "7°", "octavo": "8°", "noveno": "9°", "decimo": "10°", "once": "11°", "undecimo": "11°"}
# Escala por defecto (Decreto 1290: cada colegio fija la suya en su SIEE; se puede cambiar en la hoja «Escala»)
ESCALA = [("Superior", 4.6), ("Alto", 4.0), ("Básico", 3.0), ("Bajo", 1.0)]


def normalize(text):
    t = unicodedata.normalize("NFD", str(text).lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^\w\s°]", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def to_int(w):
    w = normalize(w)
    return int(w) if w.isdigit() else NUM_WORDS.get(w)


def safe_filename(text):
    t = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", str(text)).strip(" .")
    return re.sub(r"\s+", " ", t)[:80] or "Documento"


def _out(out_dir, sub, name, ext):
    folder = os.path.join(out_dir, sub)
    os.makedirs(folder, exist_ok=True)
    base = os.path.join(folder, safe_filename(name))
    path, n = f"{base}.{ext}", 2
    while os.path.exists(path):
        path, n = f"{base} ({n}).{ext}", n + 1
    return path


# =============================================================================
# Detectar la orden
# =============================================================================
_VERB = r"\b(?:crea|creame|haz|hazme|prepara|preparame|genera|generame|arma|armame|elabora|elaborame|disena|disename|monta|montame)\b"


def detect(text):
    """('pptx'|'notas'|'tabla', datos) o None."""
    t = normalize(text)
    if not re.search(_VERB, t):
        return None
    if re.search(r"\b(?:presentacion|presentaciones|diapositivas|power ?point|slides)\b", t):
        m = re.search(r"\b(\d+|" + "|".join(NUM_WORDS) + r")\s+(?:diapositivas|laminas|slides)\b", t)
        n = to_int(m.group(1)) if m else None
        return "pptx", {"count": max(3, min(20, n)) if n else 8, "grade": detect_grade(t), "request": text.strip()}
    if (re.search(r"\b(?:planilla|hoja|libro|excel|tabla|registro|cuadro)\b", t)
            and re.search(r"\b(?:notas|calificaciones|evaluaciones)\b", t)):
        m = re.search(r"\b(?:con\s+)?(\d+|" + "|".join(NUM_WORDS) + r")\s+(?:notas|calificaciones|columnas|evaluaciones)\b", t)
        n = to_int(m.group(1)) if m else None
        return "notas", {"count": max(1, min(15, n)) if n else 5, "request": text.strip()}
    if re.search(r"\b(?:excel|hoja de calculo|tabla)\b", t):
        return "tabla", {"request": text.strip()}
    return None


def detect_grade(t):
    m = re.search(r"\b(?:para|grado|de grado)\s+(?:el\s+)?(?:grado\s+)?(\d{1,2})\s*(?:°|grado)?\b"
                  r"(?!\s*(?:diapositivas|laminas|slides|notas|minutos))", t)
    if m and 1 <= int(m.group(1)) <= 11:
        return f"{int(m.group(1))}°"
    for w, g in GRADOS.items():
        if re.search(rf"\b(?:para|de|grado)\s+(?:el\s+)?(?:grado\s+)?{w}\b", t):
            return g
    return ""


# =============================================================================
# Presentación (PowerPoint)
# =============================================================================
def pptx_prompt(request, count, grade=""):
    return (
        "Eres un docente experto de Colombia. Prepara el contenido de una presentación para clase.\n"
        f"Petición del docente: «{request}».\n"
        + (f"Grado: {grade}. Ajusta el lenguaje y la profundidad a esa edad.\n" if grade else "")
        + f"Debe tener {count} diapositivas de contenido (sin contar la portada).\n"
        "Responde SOLO con un JSON con esta forma exacta:\n"
        '{"titulo": "…", "subtitulo": "…", "diapositivas": [\n'
        '  {"tipo": "contenido", "titulo": "…", "puntos": ["…", "…", "…"], "notas": "…"},\n'
        '  {"tipo": "pregunta", "titulo": "…", "pregunta": "…", "notas": "…"},\n'
        '  {"tipo": "actividad", "titulo": "…", "puntos": ["paso 1", "paso 2"], "notas": "…"}\n'
        "]}\n"
        "Reglas: la mayoría son de tipo «contenido», con 3 a 5 puntos cortos (máximo 14 palabras cada uno), "
        "claros y correctos. Incluye UNA diapositiva tipo «pregunta» a mitad de la presentación, con una "
        "pregunta para pensar en clase, y termina con UNA tipo «actividad» con los pasos de una actividad "
        "breve. En «notas» escribe, en 2 o 3 frases, qué decir o preguntar el docente en esa diapositiva. "
        "Español de Colombia, sin markdown ni emojis.")


def _s(v, limit=300):
    return re.sub(r"\s+", " ", str(v or "")).strip()[:limit]


def pptx_validate(d, request=""):
    if not isinstance(d, dict):
        raise ValueError("la respuesta no es un objeto")
    slides = []
    for s in d.get("diapositivas") or []:
        if not isinstance(s, dict) or not _s(s.get("titulo")):
            continue
        tipo = s.get("tipo") if s.get("tipo") in ("contenido", "pregunta", "actividad") else "contenido"
        puntos = [_s(p, 220) for p in (s.get("puntos") or []) if _s(p)][:7]
        item = {"tipo": tipo, "titulo": _s(s["titulo"], 120), "puntos": puntos, "notas": _s(s.get("notas"), 1200),
                "pregunta": _s(s.get("pregunta"), 300)}
        if tipo == "pregunta" and not item["pregunta"]:
            item["pregunta"] = puntos[0] if puntos else item["titulo"]
        if tipo != "pregunta" and not puntos:
            continue
        slides.append(item)
    if len(slides) < 2:
        raise ValueError("la presentación vino casi vacía")
    return {"titulo": _s(d.get("titulo"), 120) or _s(request, 80) or "Presentación",
            "subtitulo": _s(d.get("subtitulo"), 160), "diapositivas": slides[:25]}


# Colores (fondo claro, acento azul petróleo y naranja suave)
_C = {"dark": (0x1B, 0x2A, 0x41), "accent": (0x0F, 0x76, 0x8A), "warm": (0xF2, 0x8C, 0x28),
      "light": (0xF5, 0xF7, 0xFA), "muted": (0x5B, 0x67, 0x78), "white": (0xFF, 0xFF, 0xFF)}


def build_pptx(m, out_dir, docente="", institucion="", grade=""):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt

    def rgb(k):
        return RGBColor(*_C[k])

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height

    def bg(slide, color):
        f = slide.background.fill
        f.solid()
        f.fore_color.rgb = rgb(color)

    def box(slide, x, y, w, h, color):
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(color)
        s.line.fill.background()
        return s

    def text(slide, x, y, w, h, txt, size, color="dark", bold=False, align=None, anchor=None):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        if anchor:
            tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.text = txt
        p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(size), bold, rgb(color), "Calibri"
        if align:
            p.alignment = align
        return tf

    def footer(slide, n):
        pie = " · ".join(x for x in (institucion, docente, grade) if x)
        if pie:
            text(slide, Inches(0.6), H - Inches(0.55), Inches(10), Inches(0.4), pie, 11, "muted")
        text(slide, W - Inches(1.4), H - Inches(0.55), Inches(0.8), Inches(0.4), str(n), 11, "muted",
             align=PP_ALIGN.RIGHT)

    # Portada
    s = prs.slides.add_slide(blank)
    bg(s, "dark")
    box(s, 0, Inches(5.1), W, Inches(0.12), "warm")
    text(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(2.2), m["titulo"], 44, "white", True,
         anchor=MSO_ANCHOR.BOTTOM)
    if m["subtitulo"]:
        text(s, Inches(0.9), Inches(4.1), Inches(11.5), Inches(0.9), m["subtitulo"], 22, "light")
    datos = " · ".join(x for x in (institucion, docente, grade, datetime.date.today().strftime("%d/%m/%Y")) if x)
    text(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(0.6), datos, 16, "light")

    for i, d in enumerate(m["diapositivas"], start=2):
        s = prs.slides.add_slide(blank)
        if d["tipo"] == "pregunta":
            bg(s, "accent")
            text(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(0.8), d["titulo"].upper(), 18, "light", True)
            text(s, Inches(1.2), Inches(1.8), Inches(10.9), Inches(3.8), d["pregunta"], 40, "white", True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            bg(s, "light")
            box(s, 0, 0, Inches(0.28), H, "warm" if d["tipo"] == "actividad" else "accent")
            etiqueta = "ACTIVIDAD" if d["tipo"] == "actividad" else ""
            if etiqueta:
                text(s, Inches(0.9), Inches(0.35), Inches(6), Inches(0.4), etiqueta, 14, "warm", True)
            text(s, Inches(0.9), Inches(0.7), Inches(11.6), Inches(1.1), d["titulo"], 34, "dark", True,
                 anchor=MSO_ANCHOR.MIDDLE)
            box(s, Inches(0.9), Inches(1.85), Inches(1.2), Inches(0.07), "warm")
            tb = s.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.6), Inches(4.4))
            tf = tb.text_frame
            tf.word_wrap = True
            size = 26 if len(d["puntos"]) <= 4 else 22
            for j, pt in enumerate(d["puntos"]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                marca = f"{j + 1}.  " if d["tipo"] == "actividad" else "•  "
                r1 = p.add_run()
                r1.text = marca
                r1.font.size, r1.font.bold, r1.font.color.rgb = Pt(size), True, rgb("warm" if d["tipo"] == "actividad" else "accent")
                r2 = p.add_run()
                r2.text = pt
                r2.font.size, r2.font.color.rgb, r2.font.name = Pt(size), rgb("dark"), "Calibri"
                p.space_after = Pt(14)
        footer(s, i)
        if d["notas"]:
            s.notes_slide.notes_text_frame.text = d["notas"]

    path = _out(out_dir, "Material de clase", f"Presentación - {m['titulo']}", "pptx")
    prs.save(path)
    return path


# =============================================================================
# Planilla de notas (Excel)
# =============================================================================
def build_grades_xlsx(course, names, count, out_dir, docente="", institucion="", periodo=""):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Notas"
    esc = wb.create_sheet("Escala")
    esc["A1"], esc["B1"] = "Desempeño", "Desde (nota mínima)"
    for i, (nivel, desde) in enumerate(ESCALA, start=2):
        esc.cell(row=i, column=1, value=nivel)
        esc.cell(row=i, column=2, value=desde)
    esc["A7"] = "Cambia las notas mínimas según el SIEE de tu colegio: la hoja Notas se actualiza sola."
    esc["A7"].font = Font(italic=True, color="5B6778")
    for c in ("A1", "B1"):
        esc[c].font = Font(bold=True, color="FFFFFF")
        esc[c].fill = PatternFill("solid", fgColor="1B2A41")
    esc.column_dimensions["A"].width, esc.column_dimensions["B"].width = 16, 22

    thin = Side(style="thin", color="C9D1DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="0F768A")
    first_note, last_note = 3, 2 + count
    c_prom, c_des = last_note + 1, last_note + 2
    L = get_column_letter

    ws["A1"] = f"Planilla de notas · {course}"
    ws["A1"].font = Font(size=16, bold=True, color="1B2A41")
    info = [("Institución", institucion), ("Docente", docente), ("Periodo", periodo),
            ("Fecha", datetime.date.today().strftime("%d/%m/%Y"))]
    ws["A2"] = "   ·   ".join(f"{k}: {v}" for k, v in info if v)
    ws["A2"].font = Font(color="5B6778")
    hr = 4
    heads = ["N°", "Estudiante"] + [f"Nota {i}" for i in range(1, count + 1)] + ["Promedio", "Desempeño"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=hr, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[hr].height = 30
    names = list(names) or [""] * 30
    r0 = hr + 1
    for i, name in enumerate(names):
        r = r0 + i
        ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=name)
        rng = f"{L(first_note)}{r}:{L(last_note)}{r}"
        ws.cell(row=r, column=c_prom, value=f'=IF(COUNT({rng})=0,"",ROUND(AVERAGE({rng}),1))')
        p = f"{L(c_prom)}{r}"
        ws.cell(row=r, column=c_des, value=(
            f'=IF({p}="","",IF({p}>=Escala!$B$2,Escala!$A$2,IF({p}>=Escala!$B$3,Escala!$A$3,'
            f'IF({p}>=Escala!$B$4,Escala!$A$4,Escala!$A$5))))'))
        for j in range(1, c_des + 1):
            cell = ws.cell(row=r, column=j)
            cell.border = border
            if j >= first_note:
                cell.alignment = Alignment(horizontal="center")
            if first_note <= j <= c_prom:
                cell.number_format = "0.0"
        if i % 2:
            for j in range(1, c_des + 1):
                ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor="F5F7FA")
    rl = r0 + len(names) - 1
    # notas válidas: de 1,0 a 5,0
    dv = DataValidation(type="decimal", operator="between", formula1="1", formula2="5", allow_blank=True,
                        showErrorMessage=True, errorTitle="Nota fuera de rango",
                        error="Escribe una nota entre 1,0 y 5,0.")
    ws.add_data_validation(dv)
    dv.add(f"{L(first_note)}{r0}:{L(last_note)}{rl}")
    # colores por desempeño
    colores = {"$A$2": "C6EFCE", "$A$3": "DDEBF7", "$A$4": "FFF2CC", "$A$5": "F8CBAD"}
    rng = f"{L(c_des)}{r0}:{L(c_des)}{rl}"
    for ref, color in colores.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"{L(c_des)}{r0}=Escala!{ref}"],
                                                       fill=PatternFill("solid", fgColor=color)))
    ws.conditional_formatting.add(f"{L(first_note)}{r0}:{L(c_prom)}{rl}", FormulaRule(
        formula=[f'AND(ISNUMBER({L(first_note)}{r0}),{L(first_note)}{r0}<Escala!$B$4)'], font=Font(color="C00000", bold=True)))
    # resumen
    rs = rl + 2
    ws.cell(row=rs, column=2, value="Promedio del curso").font = Font(bold=True)
    pr = f"{L(c_prom)}{r0}:{L(c_prom)}{rl}"
    ws.cell(row=rs, column=c_prom, value=f'=IF(COUNT({pr})=0,"",ROUND(AVERAGE({pr}),1))').number_format = "0.0"
    for k, (nivel, _) in enumerate(ESCALA, start=1):
        ws.cell(row=rs + k, column=2, value=f"Estudiantes en {nivel}")
        ws.cell(row=rs + k, column=c_prom, value=f'=COUNTIF({L(c_des)}{r0}:{L(c_des)}{rl},Escala!$A${k + 1})')
    # presentación
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = max(28, min(45, max((len(n) for n in names), default=10) + 4))
    for j in range(first_note, c_des + 1):
        ws.column_dimensions[L(j)].width = 11 if j < c_des else 13
    ws.freeze_panes = ws.cell(row=r0, column=first_note)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{hr}:{hr}"
    path = _out(out_dir, "Planillas", f"Notas {course}" + (f" {periodo}" if periodo else ""), "xlsx")
    wb.save(path)
    return path


# =============================================================================
# Tabla general (Excel) hecha por la IA
# =============================================================================
def table_prompt(request):
    return ("Arma una tabla para una hoja de cálculo según esta petición de un docente en Colombia: "
            f"«{request}».\nResponde SOLO con un JSON: "
            '{"titulo": "…", "columnas": ["…", "…"], "filas": [["…", 1.5], ["…", 2]]}\n'
            "Máximo 12 columnas y 200 filas. Los números van como números (no texto), sin unidades dentro "
            "de la celda (pon la unidad en el nombre de la columna). Datos correctos; si no sabes un dato, "
            "deja la celda vacía. Español de Colombia.")


def table_validate(d, request=""):
    if not isinstance(d, dict):
        raise ValueError("la respuesta no es un objeto")
    cols = [_s(c, 60) for c in (d.get("columnas") or [])][:12]
    if not cols:
        raise ValueError("la tabla no trae columnas")
    rows = []
    for r in (d.get("filas") or [])[:200]:
        if not isinstance(r, list):
            continue
        row = []
        for v in r[:len(cols)]:
            row.append(v if isinstance(v, (int, float)) and not isinstance(v, bool) else _s(v, 500))
        rows.append(row + [""] * (len(cols) - len(row)))
    if not rows:
        raise ValueError("la tabla vino vacía")
    return {"titulo": _s(d.get("titulo"), 100) or _s(request, 80) or "Tabla", "columnas": cols, "filas": rows}


def build_table_xlsx(m, out_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = safe_filename(m["titulo"])[:31] or "Tabla"
    ws["A1"] = m["titulo"]
    ws["A1"].font = Font(size=15, bold=True, color="1B2A41")
    hr = 3
    for j, c in enumerate(m["columnas"], start=1):
        cell = ws.cell(row=hr, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F768A")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(m["filas"], start=hr + 1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    last = hr + len(m["filas"])
    ref = f"A{hr}:{get_column_letter(len(m['columnas']))}{last}"
    # nombres de columna únicos (requisito de las tablas de Excel)
    seen = {}
    for j in range(1, len(m["columnas"]) + 1):
        c = ws.cell(row=hr, column=j)
        base = str(c.value or f"Columna {j}")
        seen[base] = seen.get(base, 0) + 1
        if seen[base] > 1:
            c.value = f"{base} {seen[base]}"
    t = Table(displayName="Tabla1", ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for j, c in enumerate(m["columnas"], start=1):
        vals = [len(str(r[j - 1])) for r in m["filas"][:100]] + [len(c)]
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(50, max(vals) + 3))
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    path = _out(out_dir, "Planillas", m["titulo"], "xlsx")
    wb.save(path)
    return path
