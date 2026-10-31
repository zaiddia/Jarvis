"""Exportaciones (8.7): el quiz para Kahoot o Quizizz (Wayground) y las notas del seguimiento para Zeti.

Kahoot y Quizizz importan preguntas desde una hoja de Excel con su propio formato. Zeti (la plataforma de tu
institución) no publica un formato de importación: Jarvis arma una planilla clara por curso (actividades, definitiva
en escala 1,0–5,0, desempeño y fallas) y una hoja «Para copiar» con solo el nombre y la definitiva, en el orden de la
lista, lista para pasar a la planilla de Zeti.
"""
import datetime
import os
import re
import unicodedata

LETTERS = "abcd"
KAHOOT_TIMES = (5, 10, 20, 30, 60, 90, 120, 240)


def _safe(text):
    return re.sub(r"\s+", " ", re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", str(text))).strip()[:70] or "export"


def _clip(text, n):
    t = re.sub(r"\s+", " ", str(text)).strip()
    return t if len(t) <= n else t[:n - 1].rstrip() + "…"


def _answer_index(q):
    k = LETTERS.find(str(q.get("respuesta", "a"))[:1].lower())
    return k if 0 <= k < 4 else 0


def kahoot_rows(quiz, seconds=20):
    """Filas (pregunta, r1..r4, tiempo, correcta) con los límites de Kahoot (120 y 75 caracteres)."""
    t = min(KAHOOT_TIMES, key=lambda x: abs(x - seconds))
    rows = []
    for q in quiz.get("preguntas") or []:
        text = q["enunciado"] if not q.get("contexto") else f"{q['contexto']} {q['enunciado']}"
        opts = [_clip(o, 75) for o in q["opciones"]][:4]
        rows.append([_clip(text, 120), *opts, t, str(_answer_index(q) + 1)])
    return rows


def kahoot_xlsx(quiz, out_dir, seconds=20):
    """Plantilla de Kahoot: encabezados en la fila 8, preguntas desde la fila 9 (columna A con el número)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = "Quiz template"
    ws["B2"].font = Font(bold=True, size=14)
    ws["B3"] = f"Creado con Jarvis · {quiz.get('titulo', '')}"
    ws["B4"] = "Kahoot › Crear › Añadir pregunta › Importar › Importar hoja de cálculo"
    heads = ["Question - max 120 characters", "Answer 1 - max 75 characters", "Answer 2 - max 75 characters",
             "Answer 3 - max 75 characters", "Answer 4 - max 75 characters",
             "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs", "Correct answer(s) - choose at least one"]
    for j, h in enumerate(heads, 2):
        ws.cell(row=8, column=j, value=h).font = Font(bold=True)
    for i, r in enumerate(kahoot_rows(quiz, seconds), 1):
        ws.cell(row=8 + i, column=1, value=i)
        for j, v in enumerate(r, 2):
            ws.cell(row=8 + i, column=j, value=v)
    for col, w in zip("ABCDEFGH", (5, 60, 30, 30, 30, 30, 18, 18)):
        ws.column_dimensions[col].width = w
    return _save(wb, out_dir, f"Kahoot - {quiz.get('tema') or quiz.get('titulo', 'quiz')}")


def quizizz_rows(quiz, seconds=30):
    rows = []
    for q in quiz.get("preguntas") or []:
        text = q["enunciado"] if not q.get("contexto") else f"{q['contexto']}\n\n{q['enunciado']}"
        rows.append([text, "Multiple Choice", *q["opciones"][:4], "", _answer_index(q) + 1, seconds, "", q.get("explicacion", "")])
    return rows


def quizizz_xlsx(quiz, out_dir, seconds=30):
    """Plantilla de Quizizz (Wayground): la fila 1 con los encabezados exactos."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    heads = ["Question Text", "Question Type", "Option 1", "Option 2", "Option 3", "Option 4", "Option 5", "Correct Answer",
             "Time in seconds", "Image Link", "Answer explanation"]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in quizizz_rows(quiz, seconds):
        ws.append(r)
    ws.column_dimensions["A"].width = 60
    return _save(wb, out_dir, f"Quizizz - {quiz.get('tema') or quiz.get('titulo', 'quiz')}")


def _save(wb, out_dir, title):
    folder = os.path.join(out_dir, "Exportaciones")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{_safe(title)} {datetime.date.today():%Y-%m-%d}.xlsx")
    wb.save(path)
    return path


# ----------------------------------------------------------------------------- Zeti
def _norm(t):
    t = unicodedata.normalize("NFD", str(t).lower())
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


def surname_key(name):
    """«Juan Carlos Pérez Gómez» → ordena por «perez gomez juan carlos» (los dos últimos son los apellidos)."""
    p = str(name).split()
    if len(p) >= 4:
        return _norm(" ".join(p[-2:] + p[:-2]))
    if len(p) >= 2:
        return _norm(" ".join(p[-1:] + p[:-1]))
    return _norm(name)


def performance(avg):
    import zeti
    return zeti.performance(avg)


def zeti_table(names, grades, absences, since="", order="lista"):
    """(actividades, filas). Filas: [nombre, nota por actividad…, definitiva, desempeño, fallas]."""
    since = since or ""
    acts = []
    for n in names:
        for g in grades.get(n, []):
            if g.get("fecha", "") >= since and g["actividad"] not in acts:
                acts.append(g["actividad"])
    ordered = sorted(names, key=surname_key) if order == "apellido" else list(names)
    rows = []
    for n in ordered:
        gs = [g for g in grades.get(n, []) if g.get("fecha", "") >= since]
        by = {}
        for g in gs:
            by.setdefault(g["actividad"], []).append(g["nota"])
        cells = [round(sum(by[a]) / len(by[a]), 1) if a in by else None for a in acts]
        notas = [g["nota"] for g in gs]
        avg = round(sum(notas) / len(notas), 1) if notas else None
        rows.append([n, *cells, avg, performance(avg), absences.get(n, 0)])
    return acts, rows


def zeti_xlsx(curso, names, grades, absences, out_dir, since="", order="lista", comma=True, institucion="", docente="",
              asignatura="", weights=None):
    """10.1: la planilla con la forma de «Valoraciones Periódicas de Desempeño» de Zeti. → (ruta, estudiantes, actividades)."""
    import zeti
    folder = os.path.join(out_dir, "Exportaciones")
    os.makedirs(folder, exist_ok=True)
    title = " ".join(x for x in (curso, asignatura) if x)
    path = os.path.join(folder, f"{_safe('Zeti ' + title)} {datetime.date.today():%Y-%m-%d}.xlsx")
    n, acts, _rows = zeti.xlsx(curso, names, grades, absences, path, since, order, comma, institucion, docente, surname_key, weights,
                               asignatura)
    return path, n, acts
